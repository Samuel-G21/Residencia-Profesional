import copy
import openpyxl
from jinja2 import Template
import os

filepath = os.path.join('backend', 'templates', 'SCPM-07.xlsx')

data = {
    "clave_evento": "EVT123",
    "id_evento": "EVT123",
    "nombre_evento": "Curso de Python",
    "nombre_curso": "Curso de Python",
    "tipo_curso": "Programacion",
    "lista_participantes": [
        {"ficha": "1001", "nombre_completo": "Alice", "nombre_trabajador": "Alice", "calificacion": 90},
        {"ficha": "1002", "nombre_completo": "Bob", "nombre_trabajador": "Bob", "calificacion": 85},
        {"ficha": "1003", "nombre_completo": "Charlie", "nombre_trabajador": "Charlie", "calificacion": 100}
    ]
}
lista_participantes = data["lista_participantes"]

wb = openpyxl.load_workbook(filepath, keep_vba=False)
for ws in wb.worksheets:
    template_row_idx = None
    for row in ws.iter_rows():
        for cell in row:
            if cell.value and isinstance(cell.value, str) and '{% tr %}' in cell.value:
                template_row_idx = cell.row
                break
        if template_row_idx:
            break

    if template_row_idx:
        num_participants = len(lista_participantes)
        
        if num_participants > 1:
            ws.insert_rows(template_row_idx + 1, num_participants - 1)
            for i in range(1, num_participants):
                for col_idx in range(1, ws.max_column + 1):
                    source = ws.cell(row=template_row_idx, column=col_idx)
                    target = ws.cell(row=template_row_idx + i, column=col_idx)
                    target.value = source.value
                    if source.has_style:
                        target.font = copy.copy(source.font)
                        target.border = copy.copy(source.border)
                        target.fill = copy.copy(source.fill)
                        target.number_format = source.number_format
                        target.protection = copy.copy(source.protection)
                        target.alignment = copy.copy(source.alignment)
        elif num_participants == 0:
            ws.delete_rows(template_row_idx)

        if num_participants > 0:
            for i, participante in enumerate(lista_participantes):
                row_data = data.copy()
                row_data['participante'] = participante
                for col_idx in range(1, ws.max_column + 1):
                    cell = ws.cell(row=template_row_idx + i, column=col_idx)
                    if cell.value and isinstance(cell.value, str) and '{' in cell.value:
                        val = cell.value.replace('{% tr %}', '')
                        val = val.replace('paticipante.nombe_completo', 'participante.nombre_completo')
                        try:
                            cell.value = Template(val).render(row_data)
                        except Exception:
                            pass

    for row in ws.iter_rows():
        if template_row_idx and template_row_idx <= row[0].row < template_row_idx + len(lista_participantes):
            continue
        for cell in row:
            if cell.value and isinstance(cell.value, str) and '{' in cell.value:
                try:
                    val = cell.value.replace('{% tr %}', '')
                    cell.value = Template(val).render(data)
                except Exception:
                    pass

wb.save("test_out.xlsx")
print("Done")
