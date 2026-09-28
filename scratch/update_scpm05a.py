import openpyxl

wb = openpyxl.load_workbook("backend/templates/SCPM-05A.xlsx")
ws = wb.active

ws['C9'] = '{{ ficha }}'
ws['I9'] = '{{ nombre_completo }}'
ws['C10'] = '{{ departamento }}'
ws['I10'] = '{{ categoria }}'

wb.save("backend/templates/SCPM-05A.xlsx")
print("Template SCPM-05A.xlsx updated with Jinja tags.")
