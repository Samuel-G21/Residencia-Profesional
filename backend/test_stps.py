import copy
import openpyxl
from jinja2 import Template
import re
import io

def mock_export_scpm07():
    lista_participantes = [
        {"ficha": "123", "nombre_completo": "Juan Perez", "nombre_trabajador": "Juan Perez", "calificacion": 90},
        {"ficha": "456", "nombre_completo": "Maria Lopez", "nombre_trabajador": "Maria Lopez", "calificacion": 100}
    ]
    data = {
        "clave_evento": "E-001",
        "id_evento": "E-001",
        "nombre_evento": "Curso Seguridad",
        "nombre_curso": "Curso Seguridad",
        "tipo_curso": "Seguridad",
        "lista_participantes": lista_participantes
    }

    filepath = 'templates/SCPM-07.xlsx'
    wb = openpyxl.load_workbook(filepath, keep_vba=False)
    
    for ws in wb.worksheets:
        template_row_idx = None
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str) and '{% tr %}' in cell.value:
                    template_row_idx = cell.row
                    break
            if template_row_idx:
                break

        if template_row_idx:
            num_participants = len(lista_participantes)
            
            if num_participants > 1:
                ws.insert_rows(template_row_idx + 1, num_participants - 1)
                for i in range(1, num_participants):
                    for col_idx in range(1, ws.max_column + 1):
                        source = ws.cell(row=template_row_idx, column=col_idx)
                        target = ws.cell(row=template_row_idx + i, column=col_idx)
                        target.value = source.value
                        if source.has_style:
                            target.font = copy.copy(source.font)
                            target.border = copy.copy(source.border)
                            target.fill = copy.copy(source.fill)
                            target.number_format = source.number_format
                            target.protection = copy.copy(source.protection)
                            target.alignment = copy.copy(source.alignment)
            elif num_participants == 0:
                ws.delete_rows(template_row_idx)

            if num_participants > 0:
                for i, participante in enumerate(lista_participantes):
                    row_data = data.copy()
                    row_data['participante'] = participante
                    for col_idx in range(1, ws.max_column + 1):
                        cell = ws.cell(row=template_row_idx + i, column=col_idx)
                        if cell.value and isinstance(cell.value, str) and '{' in cell.value:
                            val = cell.value
                            # Original replacements:
                            val = val.replace('{% tr %}', '')
                            val = val.replace('paticipante.nombe_completo', 'participante.nombre_completo')
                            val = val.replace('{% for participante in lista_participantes %}', '')
                            val = val.replace('{% endfor %}', '')
                            
                            try:
                                cell.value = Template(val).render(row_data)
                            except Exception as e:
                                print(f"Template Error in cell {cell.coordinate}: '{val}' - Error: {e}")

        # Now do the outside replacements
        for row in ws.iter_rows():
            if template_row_idx and template_row_idx <= row[0].row < template_row_idx + len(lista_participantes):
                continue
            for cell in row:
                if cell.value and isinstance(cell.value, str) and '{' in cell.value:
                    try:
                        val = cell.value.replace('{% tr %}', '')
                        val = val.replace('{% for participante in lista_participantes %}', '')
                        val = val.replace('{% endfor %}', '')
                        cell.value = Template(val).render(data)
                    except Exception as e:
                        print(f"Template Error in outside cell {cell.coordinate}: '{val}' - Error: {e}")

    # Just print the result
    ws = wb.worksheets[0]
    for row in ws.iter_rows():
        for cell in row:
            if cell.value:
                print(f"{cell.coordinate}: {cell.value}")

if __name__ == '__main__':
    mock_export_scpm07()
