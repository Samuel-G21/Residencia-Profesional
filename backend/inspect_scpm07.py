import openpyxl
import os

template_path = "/app/templates/SCPM-07.xlsx"

if os.path.exists(template_path):
    wb = openpyxl.load_workbook(template_path, data_only=True)
    for sheet_name in wb.sheetnames:
        print(f"Sheet: {sheet_name}")
        ws = wb[sheet_name]
        for row in range(25, 45):
            row_data = []
            for col in range(1, 25):
                val = ws.cell(row=row, column=col).value
                if val is not None:
                    row_data.append(f"({row},{col}): {val}")
            if row_data:
                print(" | ".join(row_data))
else:
    print("Template not found!")
