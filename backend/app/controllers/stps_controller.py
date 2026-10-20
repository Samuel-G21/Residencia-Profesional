"""
Controlador de Exportacion de Formatos STPS y SCPM-07 (`stps_controller.py`).

Guia para practicantes:
Gestiona dos formatos especiales de Excel posteriores a la generacion inicial del evento:
1. `export_stps` (`POST /api/evento/<id_evento>/stps/export`):
   Toma la plantilla con macros `catalogos/STPS.xlsm`, inyecta la lista de participantes
   activos (excluyendo 'BAJA' y 'NO APTO') y sus calificaciones mediante Jinja2 en cada celda,
   guarda una copia en la carpeta del evento y la descarga al navegador.
2. `create_scpm07` / `export_scpm07` (`POST /api/evento/<id_evento>/scpm07/export`):
   Toma la plantilla `templates/SCPM-07.xlsx` (Evaluacion de la Capacitacion). Como el formato
   oficial SCPM-07 solo tiene espacio para 5 trabajadores por hoja (columnas I a M, indices 9 a 13),
   divide la lista de participantes en bloques (`chunks`) de 5 y clona la hoja original tantas
   veces como sea necesario (`SCPM-07 (1)`, `SCPM-07 (2)`, etc.).
"""
from flask import request, jsonify, send_file
from ..utils.security import error_response, event_dir, render_template_string
import os
import io
import json
import copy
import openpyxl
from ..routes.stps import stps_bp
from ..models.course import Curso
from ..models.history import HistorialCapacitacion


@stps_bp.route('/evento/<id_evento>/stps/export', methods=['POST'])
def export_stps(id_evento):
    """
    Endpoint `POST /api/evento/<id_evento>/stps/export`.
    Genera y descarga el libro `STPS_<id_evento>.xlsm` con los trabajadores activos del curso.
    """
    try:
        curso = Curso.query.filter_by(id_evento=id_evento).first()
        if not curso:
            return jsonify({"status": "error", "message": "Evento no encontrado."}), 404

        # Solo se incluyen en STPS los trabajadores que aprobaron/siguen activos (no BAJA ni NO APTO)
        trabajadores = HistorialCapacitacion.query.filter_by(id_evento=id_evento).all()
        lista_participantes = []
        for t in trabajadores:
            if t.estado not in ('BAJA', 'NO APTO'):
                lista_participantes.append({
                    "ficha": t.ficha_trabajador,
                    "nombre_completo": t.nombre_trabajador,
                    "nombre_trabajador": t.nombre_trabajador,
                    "calificacion": t.calificacion if t.calificacion is not None else ""
                })

        data = {
            "clave_evento": id_evento,
            "id_evento": id_evento,
            "nombre_evento": curso.nombre_evento,
            "nombre_curso": curso.nombre_evento,
            "tipo_curso": curso.tipo_curso,
            "lista_participantes": lista_participantes
        }

        req_data = request.get_json(silent=True)
        if req_data:
            data.update(req_data)

        filepath = os.path.join('catalogos', 'STPS.xlsm')
        if not os.path.exists(filepath):
            return jsonify({"status": "error", "message": "No se ha subido el archivo base STPS.xlsm en catalogos."}), 404

        # keep_vba=True conserva las macros VBA propias del archivo STPS (.xlsm)
        wb = openpyxl.load_workbook(filepath, keep_vba=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str) and '{{' in cell.value:
                        try:
                            cell.value = render_template_string(cell.value, data)
                            ws.row_dimensions[cell.row].height = None
                            if cell.alignment:
                                new_align = copy.copy(cell.alignment)
                                new_align.wrap_text = True
                                new_align.shrink_to_fit = False
                                cell.alignment = new_align
                            else:
                                cell.alignment = openpyxl.styles.Alignment(wrap_text=True)
                        except Exception as e:
                            print(f"Error procesando celda {cell.coordinate}: {e}")

        mem_file = io.BytesIO()
        wb.save(mem_file)

        # Guardar una copia en la carpeta de salida del evento para incluirla en el ZIP consolidado
        evento_dir = event_dir(str(id_evento).strip(), create=True)
        os.makedirs(evento_dir, exist_ok=True)
        out_name = f'STPS_{id_evento}.xlsm'
        with open(os.path.join(evento_dir, out_name), 'wb') as f:
            f.write(mem_file.getvalue())

        mem_file.seek(0)

        return send_file(
            mem_file,
            mimetype='application/vnd.ms-excel.sheet.macroEnabled.12',
            as_attachment=True,
            download_name=f'STPS_{id_evento}.xlsm'
        )

    except Exception as e:
        return error_response("Ocurrió un error interno", 500, e)


def create_scpm07(id_evento, req_data=None):
    """
    Genera el archivo de evaluacion `SCPM-07` para un evento, paginando de 5 en 5 participantes
    por hoja de calculo y persistiendo el resultado en `outputs/<id_evento>/`.
    """
    curso = Curso.query.filter_by(id_evento=id_evento).first()
    if not curso:
        raise Exception("Evento no encontrado.")

    trabajadores = HistorialCapacitacion.query.filter_by(id_evento=id_evento).all()
    lista_participantes = []
    for t in trabajadores:
        if t.estado not in ('BAJA', 'NO APTO'):
            lista_participantes.append({
                "ficha": t.ficha_trabajador,
                "nombre_completo": t.nombre_trabajador,
                "nombre_trabajador": t.nombre_trabajador,
                "calificacion": t.calificacion if t.calificacion is not None else ""
            })

    data = {
        "clave_evento": id_evento,
        "id_evento": id_evento,
        "nombre_evento": curso.nombre_evento,
        "nombre_curso": curso.nombre_evento,
        "tipo_curso": curso.tipo_curso,
        "lista_participantes": lista_participantes
    }

    # Recupera metadatos adicionales del evento (fechas, instructor, supervisor) desde event_context.json
    evento_dir = event_dir(str(id_evento).strip(), create=True)
    ctx_path = os.path.join(evento_dir, 'event_context.json')
    if os.path.exists(ctx_path):
        with open(ctx_path, 'r', encoding='utf-8') as f:
            event_ctx = json.load(f)
            data.update(event_ctx)

    if req_data:
        data.update(req_data)

    filepath = os.path.join('templates', 'SCPM-07.xlsx')
    if not os.path.exists(filepath):
        raise Exception("No se ha subido el archivo base SCPM-07.xlsx en templates.")

    wb = openpyxl.load_workbook(filepath, keep_vba=False)
    original_ws = wb.worksheets[0]
    original_title = original_ws.title

    # El formato SCPM-07 aloja hasta 5 participantes por hoja (columnas I..M)
    chunk_size = 5
    participantes_chunks = [
        lista_participantes[i:i + chunk_size]
        for i in range(0, len(lista_participantes), chunk_size)
    ]
    if not participantes_chunks:
        participantes_chunks = [[]]

    worksheets = []
    for chunk_idx in range(len(participantes_chunks)):
        ws = wb.copy_worksheet(original_ws)
        if len(participantes_chunks) > 1:
            ws.title = f"SCPM-07 ({chunk_idx + 1})"
        else:
            ws.title = original_title + " (new)"
        worksheets.append(ws)

    wb.remove(original_ws)
    if len(participantes_chunks) == 1:
        worksheets[0].title = original_title

    for chunk_idx, chunk in enumerate(participantes_chunks):
        ws = worksheets[chunk_idx]

        ficha_row = None
        nombre_row = None

        # Localiza las filas donde se inyectan las fichas y los nombres de los 5 participantes
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    if '{% for participante in lista_participantes %}' in cell.value or 'participante.ficha' in cell.value:
                        ficha_row = cell.row
                    if 'paticipante.nombe_completo' in cell.value or 'participante.nombre_completo' in cell.value:
                        nombre_row = cell.row

        if ficha_row and nombre_row:
            # Limpia las 5 columnas de participantes (columnas 9 a 13: I, J, K, L, M)
            for col in range(9, 14):
                ws.cell(row=ficha_row, column=col).value = ""
                ws.cell(row=nombre_row, column=col).value = ""

            # Escribe cada participante del bloque actual en su columna correspondiente
            for i, p in enumerate(chunk):
                col = 9 + i
                ws.cell(row=ficha_row, column=col).value = p.get('ficha', '')
                ws.cell(row=nombre_row, column=col).value = p.get('nombre_completo', '')

        # Renderiza el resto de etiquetas Jinja de la hoja (encabezados del evento, fechas, instructor, etc.)
        for row in ws.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str) and '{' in cell.value:
                    val = cell.value.replace('{% tr %}', '')
                    val = val.replace('paticipante.nombe_completo', 'participante.nombre_completo')
                    val = val.replace('{% for participante in lista_participantes %}', '')
                    val = val.replace('{% endfor %}', '')
                    try:
                        cell.value = render_template_string(val, data)
                        ws.row_dimensions[cell.row].height = None
                        if cell.alignment:
                            new_align = copy.copy(cell.alignment)
                            new_align.wrap_text = True
                            new_align.shrink_to_fit = False
                            cell.alignment = new_align
                        else:
                            cell.alignment = openpyxl.styles.Alignment(wrap_text=True)
                    except Exception as e:
                        print(f"Template Error in cell {cell.coordinate}: '{cell.value}' - Error: {e}")

    mem_file = io.BytesIO()
    wb.save(mem_file)

    evento_dir = event_dir(str(id_evento).strip(), create=True)
    os.makedirs(evento_dir, exist_ok=True)
    out_name = f'{id_evento}_EVENTO_SCPM-07.xlsx'

    # Si existia un archivo con nomenclatura anterior, se elimina para evitar duplicados en el ZIP
    old_out_name = f'SCPM-07_{id_evento}.xlsx'
    old_path = os.path.join(evento_dir, old_out_name)
    if os.path.exists(old_path):
        try:
            os.remove(old_path)
        except Exception:
            pass

    with open(os.path.join(evento_dir, out_name), 'wb') as f:
        f.write(mem_file.getvalue())

    mem_file.seek(0)
    return mem_file, out_name


@stps_bp.route('/evento/<id_evento>/scpm07/export', methods=['POST'])
def export_scpm07(id_evento):
    """
    Endpoint `POST /api/evento/<id_evento>/scpm07/export`.
    Genera y descarga el archivo Excel `SCPM-07` del evento solicitado.
    """
    try:
        req_data = request.get_json(silent=True)
        mem_file, out_name = create_scpm07(id_evento, req_data)
        return send_file(
            mem_file,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=out_name
        )
    except Exception as e:
        return error_response("Ocurrió un error interno", 500, e)
