"""
Controlador de Gestion de Eventos y Participantes (`event_controller.py`).

Guia para practicantes:
Concentra las operaciones CRUD y de seguimiento sobre cursos registrados (`Curso`) y sus
participantes (`HistorialCapacitacion`):
- Consulta del listado historico de eventos y su avance por fase (1 a 5).
- Cierre o cancelacion de eventos (`FINALIZADO`, `CANCELADO`). Una vez que un evento entra en
  `ESTADOS_CERRADOS`, ya no se permite alterar su fase ni modificar a sus trabajadores.
- Alta manual de un trabajador adicional (`agregar_trabajador`), que busca su CURP/nombre en
  `catalogos/CURP.xlsx`, genera sus formatos individuales y regenera el `SCPM-07` grupal.
- Baja (`BAJA`) o reprobacion (`NO APTO`) de participantes individuales.
- Carga del formato `SCPM-07` lleno con calificaciones finales (`upload_scpm07`) y del `SCPM-03`
  firmado (`upload_scpm03`).
"""
from flask import jsonify, request
from ..utils.security import error_response, event_dir, ESTADOS_CERRADOS
from ..models.course import Curso
from ..models.history import HistorialCapacitacion
from ..models import db
from sqlalchemy import text
import pandas as pd
import os


def get_all_cursos():
    """
    Endpoint `GET /api/cursos`.
    Devuelve todos los eventos registrados ordenados del mas reciente al mas antiguo,
    incluyendo el conteo total de participantes inscritos en cada uno.
    """
    sql = text("""
        SELECT c.id_evento, c.nombre_evento, c.fase_actual, c.fecha_registro, c.estado,
               COUNT(h.id) as total_participantes
        FROM cursos c
        LEFT JOIN historial_capacitacion h ON c.id_evento = h.id_evento
        GROUP BY c.id_evento, c.nombre_evento, c.fase_actual, c.fecha_registro, c.estado
        ORDER BY c.fecha_registro DESC
    """)
    result = db.session.execute(sql)
    data = []
    for row in result:
        data.append({
            "id_evento": row.id_evento,
            "nombre_evento": row.nombre_evento,
            "fase_actual": row.fase_actual,
            "fecha_registro": row.fecha_registro.strftime('%Y-%m-%d %H:%M:%S') if row.fecha_registro else None,
            "estado": row.estado,
            "total_participantes": row.total_participantes
        })
    return jsonify({"status": "success", "data": data}), 200


def get_evento_fase(id_evento):
    """
    Endpoint `GET /api/evento/<id_evento>`.
    Consulta la fase actual y estado de un evento por su ID. Si no existe aun en la BD,
    responde `not_found` sugiriendo iniciar en la Fase 1.
    """
    curso = Curso.query.filter_by(id_evento=id_evento).first()
    if curso:
        return jsonify({
            "status": "success",
            "data": {
                "fase_actual": curso.fase_actual,
                "nombre_evento": curso.nombre_evento,
                "estado": curso.estado,
                "id_evento": curso.id_evento
            }
        }), 200
    return jsonify({"status": "not_found", "message": "Evento nuevo", "fase_sugerida": 1}), 200


def get_evento_trabajadores(id_evento):
    """
    Endpoint `GET /api/evento/<id_evento>/trabajadores`.
    Lista todos los trabajadores inscritos en un evento junto con su estado y calificacion.
    """
    sql = text(
        "SELECT id, ficha_trabajador, nombre_trabajador, estado, calificacion "
        "FROM historial_capacitacion WHERE id_evento = :id"
    )
    result = db.session.execute(sql, {"id": id_evento})
    data = []
    for row in result:
        data.append({
            "id": row.id,
            "ficha_trabajador": row.ficha_trabajador,
            "nombre_trabajador": row.nombre_trabajador,
            "estado": row.estado,
            "calificacion": float(row.calificacion) if row.calificacion is not None else None
        })
    return jsonify({"status": "success", "data": data}), 200


def update_evento_fase(id_evento):
    """
    Endpoint `PUT /api/evento/<id_evento>/fase`.
    Actualiza la fase actual del evento (1 a 5), siempre que no se encuentre finalizado o cancelado.
    """
    data = request.json
    nueva_fase = data.get('fase')
    if not nueva_fase:
        return jsonify({"status": "error", "message": "Falta la nueva fase"}), 400

    curso = Curso.query.filter_by(id_evento=id_evento).first()
    if curso and curso.estado in ESTADOS_CERRADOS:
        return jsonify({"status": "error", "message": f"No se puede modificar un evento en estado {curso.estado}"}), 400

    if curso:
        curso.fase_actual = nueva_fase
    else:
        curso = Curso(id_evento=id_evento, nombre_evento='Evento Manual', fase_actual=nueva_fase)
        db.session.add(curso)
    db.session.commit()
    return jsonify({"status": "success", "message": f"Fase actualizada a {nueva_fase}"}), 200


def delete_evento(id_evento):
    """
    Endpoint `DELETE /api/evento/<id_evento>`.
    Elimina un curso de la base de datos (y en cascada su historial de participantes).
    """
    curso = Curso.query.filter_by(id_evento=id_evento).first()
    if curso:
        db.session.delete(curso)
        db.session.commit()
        return jsonify({"status": "success", "message": "Evento eliminado"}), 200
    return jsonify({"status": "error", "message": "Evento no encontrado"}), 404


def finalizar_evento(id_evento):
    """
    Endpoint `PUT /api/evento/<id_evento>/finalizar`.
    Marca un evento como `FINALIZADO`, bloqueando modificaciones posteriores.
    """
    curso = Curso.query.filter_by(id_evento=id_evento).first()
    if curso:
        curso.estado = 'FINALIZADO'
        db.session.commit()
        return jsonify({"status": "success", "message": "Evento finalizado"}), 200
    return jsonify({"status": "error", "message": "Evento no encontrado"}), 404


def cancelar_evento(id_evento):
    """
    Endpoint `PUT /api/evento/<id_evento>/cancelar`.
    Marca un evento como `CANCELADO`, bloqueando modificaciones posteriores.
    """
    curso = Curso.query.filter_by(id_evento=id_evento).first()
    if curso:
        curso.estado = 'CANCELADO'
        db.session.commit()
        return jsonify({"status": "success", "message": "Evento cancelado"}), 200
    return jsonify({"status": "error", "message": "Evento no encontrado"}), 404


def baja_trabajador(id_evento, ficha):
    """
    Endpoint `PUT /api/evento/<id_evento>/trabajador/<ficha>/baja`.
    Cambia el estado de un trabajador a `BAJA` y guarda el motivo correspondiente.
    """
    data = request.json or {}
    motivo = data.get('motivo', '')
    curso = Curso.query.filter_by(id_evento=id_evento).first()
    if curso and curso.estado in ESTADOS_CERRADOS:
        return jsonify({"status": "error", "message": f"No se puede modificar trabajadores de un evento en estado {curso.estado}"}), 400

    trabajador = HistorialCapacitacion.query.filter_by(id_evento=id_evento, ficha_trabajador=ficha).first()
    if trabajador:
        trabajador.estado = 'BAJA'
        trabajador.motivo_baja = motivo
        db.session.commit()
        return jsonify({"status": "success", "message": "Trabajador dado de baja"}), 200
    return jsonify({"status": "error", "message": "Trabajador no encontrado en el evento"}), 404


def no_apto_trabajador(id_evento, ficha):
    """
    Endpoint `PUT /api/evento/<id_evento>/trabajador/<ficha>/no_apto`.
    Marca a un trabajador como `NO APTO` y registra su calificacion reprobatoria.
    """
    data = request.json or {}
    calificacion = data.get('calificacion')
    curso = Curso.query.filter_by(id_evento=id_evento).first()
    if curso and curso.estado in ESTADOS_CERRADOS:
        return jsonify({"status": "error", "message": f"No se puede modificar trabajadores de un evento en estado {curso.estado}"}), 400

    trabajador = HistorialCapacitacion.query.filter_by(id_evento=id_evento, ficha_trabajador=ficha).first()
    if trabajador:
        trabajador.estado = 'NO APTO'
        try:
            trabajador.calificacion = float(calificacion) if calificacion is not None else None
        except ValueError:
            pass
        db.session.commit()
        return jsonify({"status": "success", "message": "Trabajador marcado como No Apto"}), 200
    return jsonify({"status": "error", "message": "Trabajador no encontrado en el evento"}), 404


def agregar_trabajador(id_evento):
    """
    Endpoint `POST /api/evento/<id_evento>/trabajador`.
    Inscribe manualmente a un nuevo trabajador en un evento existente:
    1. Si no se envio nombre o CURP, intenta autocompletarlos buscando la ficha en `catalogos/CURP.xlsx`.
    2. Registra al trabajador en `historial_capacitacion`.
    3. Genera sus documentos individuales y actualiza el formato grupal `SCPM-07`.
    """
    curso = Curso.query.filter_by(id_evento=id_evento).first()
    if curso and curso.estado in ESTADOS_CERRADOS:
        return jsonify({"status": "error", "message": f"No se puede modificar trabajadores de un evento en estado {curso.estado}"}), 400

    data = request.json or {}
    ficha = str(data.get('ficha', '')).strip()
    nombre = str(data.get('nombre', '')).strip()

    if not ficha:
        return jsonify({"status": "error", "message": "Falta la ficha"}), 400

    # Autocompletado opcional desde el catalogo maestro CURP.xlsx
    if not nombre or not data.get('curp'):
        ruta_curp = os.path.join('catalogos', 'CURP.xlsx')
        if os.path.exists(ruta_curp):
            try:
                df_curp = pd.read_excel(ruta_curp)
                df_curp.columns = df_curp.columns.str.upper().str.strip()
                match = df_curp[df_curp['FICHA'].astype(str) == ficha]
                if not match.empty:
                    if not nombre and 'NOMBRE' in df_curp.columns:
                        nombre = str(match.iloc[0]['NOMBRE']).strip().title()
                    if not data.get('curp') and 'CURP' in df_curp.columns:
                        data['curp'] = str(match.iloc[0]['CURP']).upper()
            except Exception as e:
                print("Error buscando en catalogo:", e)

    if not nombre:
        return jsonify({
            "status": "error",
            "message": "No se encontró el nombre en el catálogo. Por favor, ingrésalo manualmente."
        }), 400

    data['nombre'] = nombre

    existe = HistorialCapacitacion.query.filter_by(id_evento=id_evento, ficha_trabajador=ficha).first()
    if existe:
        return jsonify({"status": "error", "message": "El trabajador ya existe en este evento"}), 400

    nuevo = HistorialCapacitacion(
        id_evento=id_evento,
        ficha_trabajador=ficha,
        nombre_trabajador=nombre,
        estado='ACTIVO'
    )
    db.session.add(nuevo)
    db.session.commit()

    # Genera los formatos individuales del nuevo participante y regenera el SCPM-07 del evento
    try:
        from .document_controller import generate_single_worker_docs
        generate_single_worker_docs(id_evento, data)

        from .stps_controller import create_scpm07
        try:
            create_scpm07(id_evento)
        except Exception as scpm_e:
            print("Error regenerando SCPM-07:", scpm_e)

    except Exception as e:
        print("Error generating docs for new worker:", e)

    return jsonify({"status": "success", "message": "Trabajador agregado correctamente"}), 200


def upload_scpm07(id_evento):
    """
    Endpoint `POST /api/evento/<id_evento>/scpm07`.
    Extrae las calificaciones de los participantes desde un archivo Excel SCPM-07 subido:
    - Soporta el formato oficial con la hoja `'Propuesta'` (fichas en fila 30, promedio en fila ~42).
    - Soporta tambien una tabla plana con columnas `FICHA` y `CALIFICACION`.
    Guarda una copia del archivo en `outputs/<id_evento>/SCPM-07_<id_evento>_Firmado.xlsx`.
    """
    curso = Curso.query.filter_by(id_evento=id_evento).first()
    if curso and curso.estado in ESTADOS_CERRADOS:
        return jsonify({"status": "error", "message": f"No se puede modificar un evento en estado {curso.estado}"}), 400

    if 'file' not in request.files:
        return jsonify({"status": "error", "message": "Falta el archivo SCPM-07"}), 400
    file = request.files['file']
    if not file.filename.endswith('.xlsx'):
        return jsonify({"status": "error", "message": "Debe ser un archivo Excel (.xlsx)"}), 400

    try:
        import openpyxl
        wb = openpyxl.load_workbook(file, data_only=True)
        # Caso 1: Formato institucional SCPM-07 con pestana 'Propuesta'
        if 'Propuesta' in wb.sheetnames:
            ws = wb['Propuesta']
            ficha_row, calif_row = 30, 42
            for r in range(25, 45):
                val = ws.cell(row=r, column=1).value
                if val and "Promedio final" in str(val):
                    calif_row = r
                    break

            for col in range(9, 30):
                ficha_val = ws.cell(row=ficha_row, column=col).value
                calif_val = ws.cell(row=calif_row, column=col).value
                if ficha_val and str(ficha_val).strip():
                    ficha_str = str(ficha_val).replace('.0', '').strip()
                    try:
                        calif = float(calif_val) if calif_val is not None else 0
                        trabajador = HistorialCapacitacion.query.filter_by(
                            id_evento=id_evento,
                            ficha_trabajador=ficha_str
                        ).first()
                        if trabajador:
                            trabajador.calificacion = calif
                    except ValueError:
                        continue
        else:
            # Caso 2: Respaldo mediante lectura tabular plana con Pandas
            file.seek(0)
            df = pd.read_excel(file)
            df.columns = df.columns.astype(str).str.upper().str.strip()
            col_ficha = next((c for c in df.columns if 'FICHA' in c), None)
            col_calif = next((c for c in df.columns if 'CALIFICACI' in c), None)

            if not col_ficha or not col_calif:
                return jsonify({
                    "status": "error",
                    "message": "No se encontró el formato oficial ni las columnas de FICHA o CALIFICACION"
                }), 400

            for _, row in df.iterrows():
                ficha = str(row.get(col_ficha, '')).replace('.0', '').strip()
                calif_raw = row.get(col_calif, '')
                if not ficha or pd.isna(calif_raw) or str(calif_raw).strip() == '':
                    continue
                try:
                    calif = float(calif_raw)
                    trabajador = HistorialCapacitacion.query.filter_by(
                        id_evento=id_evento,
                        ficha_trabajador=ficha
                    ).first()
                    if trabajador:
                        trabajador.calificacion = calif
                except ValueError:
                    continue

        db.session.commit()

        # Almacena respaldo del SCPM-07 firmado en la carpeta del expediente
        evento_dir = event_dir(str(id_evento).strip(), create=True)
        os.makedirs(evento_dir, exist_ok=True)
        out_name = f"SCPM-07_{id_evento}_Firmado.xlsx"
        file.seek(0)
        with open(os.path.join(evento_dir, out_name), 'wb') as f:
            f.write(file.read())

        return jsonify({"status": "success", "message": "Calificaciones actualizadas"}), 200
    except Exception as e:
        return error_response("Error al procesar el Excel: ", 400, e)


def upload_scpm03(id_evento):
    """
    Endpoint `POST /api/evento/<id_evento>/scpm03`.
    Recibe el formato `SCPM-03` (.docx) firmado, lo respalda en la carpeta del evento
    y avanza automaticamente el curso a la Fase 2 (Ejecucion del Evento).
    """
    curso = Curso.query.filter_by(id_evento=id_evento).first()
    if curso and curso.estado in ESTADOS_CERRADOS:
        return jsonify({"status": "error", "message": f"No se puede modificar un evento en estado {curso.estado}"}), 400

    if 'file' not in request.files:
        return jsonify({"status": "error", "message": "Falta el archivo SCPM-03"}), 400
    file = request.files['file']
    if not file.filename.endswith('.docx'):
        return jsonify({"status": "error", "message": "Debe ser un archivo Word (.docx)"}), 400

    try:
        if not curso:
            return jsonify({"status": "error", "message": "Evento no encontrado"}), 404

        evento_dir = event_dir(str(id_evento).strip(), create=True)
        os.makedirs(evento_dir, exist_ok=True)
        out_name = f"SCPM-03_{id_evento}_Firmado.docx"
        file.seek(0)
        with open(os.path.join(evento_dir, out_name), 'wb') as f:
            f.write(file.read())

        curso.fase_actual = 2
        db.session.commit()
        return jsonify({"status": "success", "message": "SCPM-03 subido y evento avanzado a Fase 2"}), 200
    except Exception as e:
        return error_response("Error al procesar el archivo: ", 400, e)
