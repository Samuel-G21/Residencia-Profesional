"""
Controlador Principal del Motor Documental (`document_controller.py`).

Guia para practicantes:
Este modulo implementa el flujo central de automatizacion documental de capacitacion:
1. `extract_pdf` (`POST /api/extract-pdf`):
   Lee el PDF oficial de asistencia (`SCPM-01`) con `pdfplumber`, extrae mediante expresiones
   regulares y tablas los metadatos del curso (ID SIRHN, nombre del evento, fechas, duracion,
   instructor, supervisor) y el listado de trabajadores, devolviendo una Guia de Asistencia en `.xlsx`.
2. `generate_docs` (`POST /api/generate-docs`):
   Toma el archivo Excel validado, cruza las fichas contra `catalogos/CURP.xlsx`, extrae
   opcionalmente el supervisor desde un `SCPM-03.docx` lleno, y rellena tanto las plantillas
   individuales (DC-3, Cedula, SCPM-04, SCPM-05A, SCPM-06) como las grupales (Carta Compromiso,
   FVC, Informe Tecnico, SCPM-03, SCPM-05, SCPM-07, SIRCE_Automatizado). Guarda todo en
   `outputs/<id_evento>/` y sincroniza el curso y participantes en MySQL.
3. `download_docs` (`GET /api/download-docs/<id_evento>`):
   Envia todos los `.docx` generados al microservicio Gotenberg para convertirlos a PDF, los une
   en un `<id_evento>_PDF_Maestro.pdf` con `pypdf`, y empaqueta todos los archivos organizados
   por carpetas dentro de `Expediente_<id_evento>.zip`.
4. `generate_single_worker_docs`:
   Funcion auxiliar invocada cuando se da de alta manualmente a un trabajador adicional desde
   el detalle de un evento (`event_controller.agregar_trabajador`).
"""
from flask import request, jsonify, send_file
from ..utils.security import error_response, event_dir, render_template_string
import os
import io
import json
import copy
import re
import shutil
import zipfile
import requests
import openpyxl
import pandas as pd
import pdfplumber
from docxtpl import DocxTemplate
from pypdf import PdfWriter
from ..models.course import Curso
from ..models.history import HistorialCapacitacion
from ..models import db


def safe_str(val):
    """
    Convierte un valor de celda de Pandas/Excel a cadena limpia sin sufijos `.0`.
    Ejemplo: `333023.0` -> `'333023'`, `NaN` -> `''`.
    """
    if pd.isna(val) or val == '':
        return ""
    try:
        f_val = float(val)
        if f_val.is_integer():
            return str(int(f_val))
        return str(val).strip()
    except Exception:
        return str(val).strip()


def generate_docs():
    """
    Endpoint `POST /api/generate-docs`.
    Genera masivamente los documentos individuales y grupales de un evento a partir del Excel
    de asistencia y registra el avance en la base de datos (Fase 4: Integracion de Expediente).
    """
    if 'file' not in request.files or 'id_evento' not in request.form:
        return jsonify({"status": "error", "message": "Faltan datos"}), 400

    file = request.files['file']
    id_evento = str(request.form['id_evento']).strip()

    try:
        # 1. Lectura y filtrado del Excel de asistencia por ID del evento
        df = pd.read_excel(file)
        df.columns = df.columns.str.upper().str.strip()
        df = df.fillna('')
        if 'ID' not in df.columns:
            return jsonify({"status": "error", "message": "Falta columna ID"}), 400
        df['ID'] = df['ID'].astype(str)
        trabajadores_curso = df[df['ID'] == id_evento]

        if trabajadores_curso.empty:
            return jsonify({"status": "error", "message": "Sin trabajadores"}), 404

        # 2. Carga del catalogo maestro CURP.xlsx en un diccionario en memoria {ficha: curp}
        diccionario_curps = {}
        ruta_curp = os.path.join('catalogos', 'CURP.xlsx')
        if os.path.exists(ruta_curp):
            try:
                df_curp = pd.read_excel(ruta_curp)
                df_curp.columns = df_curp.columns.str.upper().str.strip()
                if 'FICHA' in df_curp.columns and 'CURP' in df_curp.columns:
                    for _, r_curp in df_curp.iterrows():
                        diccionario_curps[safe_str(r_curp['FICHA'])] = str(r_curp['CURP']).upper()
            except Exception as e:
                print("Error leyendo CURP.xlsx:", e)

        # 3. Extraccion de metadatos generales del evento desde la primera fila del curso
        row0 = trabajadores_curso.iloc[0]
        nombre_curso = str(row0.get('NOMBRE DEL EVENTO', '')).strip().upper()
        duracion_curso = str(row0.get('DURACION HORAS', row0.get('DURACION', ''))).strip()
        tipo_curso = request.form.get('tipo_curso', 'Actualización')

        nombre_instructor = str(row0.get('INSTRUCTOR', row0.get('NOMBRE INSTRUCTOR', ''))).strip().title()
        ficha_instructor = str(row0.get('FICHA INSTRUCTOR', '')).strip()

        nombre_supervisor = str(
            row0.get('NOMBRE VALIDA (AREA USUARIA)', row0.get('NOMBRE SUPERVISOR', row0.get('SUPERVISOR', '')))
        ).strip().title()
        ficha_supervisor = str(
            row0.get('FICHA VALIDA (AREA USUARIA)', row0.get('FICHA SUPERVISOR', ''))
        ).strip()

        # 4. Si se adjunto opcionalmente un SCPM-03.docx lleno, se extrae de sus tablas el Supervisor Tecnico
        file_scpm03 = request.files.get('file_scpm03')
        if file_scpm03 and file_scpm03.filename.endswith('.docx'):
            try:
                import docx
                doc_scpm03 = docx.Document(file_scpm03)
                for table in doc_scpm03.tables:
                    for row in table.rows:
                        row_texts = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                        if any("supervisor" in t.lower() and "evento" in t.lower() for t in row_texts):
                            for t in row_texts:
                                t_lower = t.lower()
                                if "supervisor" in t_lower and "evento" in t_lower:
                                    continue
                                if "nombre del supervisor" in t_lower:
                                    continue
                                if t.strip():
                                    ficha_match = re.search(r'\b\d{5,6}\b', t)
                                    if ficha_match:
                                        ficha_supervisor = ficha_match.group(0)
                                        name_part = t[:ficha_match.start()]
                                    else:
                                        name_part = t

                                    name_part = re.sub(r'(?i)\bF\s*-?\s*$', '', name_part)
                                    name_part = re.sub(r'(?i)\bFicha\b\s*:?', '', name_part)
                                    name_part = re.sub(r'(?i)^(Ing\.|Lic\.|Mtra\.|Mtro\.|Dr\.|Dra\.)\s*', '', name_part)
                                    name_part = re.sub(r'[^A-Za-zÑñÁÉÍÓÚáéíóú\s]', '', name_part).strip()

                                    if name_part:
                                        nombre_supervisor = name_part.title()
                                    break
            except Exception as e:
                print("Error extracting from SCPM-03:", e)

        # 5. Formateo de fechas de inicio y termino (DD/MM/AAAA)
        dia_i = safe_str(row0.get('DIA I', '')).zfill(2)
        mes_i = safe_str(row0.get('MES I', '')).zfill(2)
        anio_i = safe_str(row0.get('AÑO I', '')).split('.')[0]
        dia_t = safe_str(row0.get('DIA T', '')).zfill(2)
        mes_t = safe_str(row0.get('MES T', '')).zfill(2)
        anio_t = safe_str(row0.get('AÑO T', '')).split('.')[0]

        f_inicio = f"{dia_i}/{mes_i}/{anio_i}" if anio_i else ""
        f_termino = f"{dia_t}/{mes_t}/{anio_t}" if anio_t else ""

        # 6. Clasificacion de plantillas en individuales (1 por trabajador) y grupales (1 por evento)
        docs_seleccionados_str = request.form.get('docs_seleccionados')

        todas_ind = [
            '1. Cédula registro actualizado 2025 COMBIANADA.docx',
            '2. Constancias de Habilidades DC-3 2026 COMBINADA.docx',
            'SCPM-04 COMBINADA.docx',
            'SCPM-04.docx',
            'SCPM-06 COMBINADA.docx',
            'SCPM-05A.xlsx'
        ]

        todas_grp = [
            '5. Carta Compromiso Instructor 2026 COMBINADA.docx',
            'FVC.docx',
            'Informe Técnico Instructor 2025.docx',
            'SCPM-05 2025.docx',
            'SCPM-03.docx',
            'SCPM-07.xlsx',
            'SIRCE_Automatizado.xlsm'
        ]

        if docs_seleccionados_str:
            clean_selections = re.sub(r'[^a-z0-9]', '', docs_seleccionados_str.lower())

            def is_selected(t_name):
                c_t = re.sub(r'[^a-z0-9]', '', t_name.lower())
                if 'registro' in c_t and 'registro' in clean_selections:
                    return True
                if 'dc3' in c_t and 'dc3' in clean_selections:
                    return True
                if 'scpm07' in c_t and 'scpm07' in clean_selections:
                    return True
                if 'sirce' in c_t and 'sirce' in clean_selections:
                    return True
                return c_t in clean_selections

            plantillas_ind = [p for p in todas_ind if is_selected(p)]
            plantillas_grp = [p for p in todas_grp if is_selected(p)]
        else:
            plantillas_ind = todas_ind
            plantillas_grp = todas_grp

        # Regla institucional: SCPM-05A aplica solo para cursos de Ascenso; en Actualizacion se omite
        if tipo_curso.lower() == 'ascenso':
            pass
        else:
            plantillas_ind = [p for p in plantillas_ind if 'SCPM-05A' not in p]
            plantillas_grp = [p for p in plantillas_grp if 'SCPM-05A' not in p]

        # 7. Preparacion del directorio de salida del evento y persistencia del contexto (event_context.json)
        evento_dir = event_dir(id_evento, create=True)
        os.makedirs(evento_dir, exist_ok=True)
        for f in os.listdir(evento_dir):
            try:
                os.remove(os.path.join(evento_dir, f))
            except Exception:
                pass

        event_ctx = {
            "nombre_curso": nombre_curso,
            "duracion": duracion_curso,
            "nombre_instructor": nombre_instructor,
            "ficha_instructor": ficha_instructor,
            "fecha_inicio": f_inicio,
            "fecha_termino": f_termino,
            "dia_inicio": dia_i,
            "mes_inicio": mes_i,
            "anio_inicio": anio_i,
            "dia_termino": dia_t,
            "mes_termino": mes_t,
            "anio_termino": anio_t,
            "nombre_supervisor": nombre_supervisor,
            "ficha_supervisor": ficha_supervisor,
            "tipo_curso": tipo_curso,
            "plantillas_ind": plantillas_ind
        }
        with open(os.path.join(evento_dir, 'event_context.json'), 'w', encoding='utf-8') as fctx:
            json.dump(event_ctx, fctx, ensure_ascii=False)

        docs_gen = []
        historial = []
        lista_participantes = []

        # 8. Generacion de formatos individuales por cada trabajador del curso
        for _, row in trabajadores_curso.iterrows():
            ficha = safe_str(row.get('FICHA', ''))
            nombres = str(row.get('NOMBRE', row.get('NOMBRE(S)', ''))).strip().title()
            pat = str(row.get('PRIMER APELLIDO', '')).strip().title()
            mat = str(row.get('SEGUNDO APELLIDO', '')).strip().title()
            nombre_completo = f"{pat} {mat} {nombres}".strip()
            cat = str(row.get('CATEGORIA', row.get('CATEGORÍA', ''))).strip()
            niv = str(row.get('NIVEL', '')).strip()
            depto = str(row.get('DEPARTAMENTO', '')).strip()
            curp = diccionario_curps.get(ficha, "")

            lista_participantes.append({
                "ficha": ficha,
                "nombre_completo": nombre_completo,
                "categoria": cat,
                "nivel": niv,
                "departamento": depto
            })
            if ficha:
                historial.append({
                    "id_evento": id_evento,
                    "ficha_trabajador": ficha,
                    "nombre_trabajador": nombre_completo
                })

            ctx_ind = {
                "ficha": ficha,
                "ficha_trabajador": ficha,
                "apellido_paterno": pat,
                "apellido_materno": mat,
                "nombre": nombres,
                "nombres": nombres,
                "nombre_completo": nombre_completo,
                "nombre_trabajador": nombre_completo,
                "curp": curp,
                "CURP": curp,
                "nombre_evento": nombre_curso,
                "nombre_curso": nombre_curso,
                "clave_evento": id_evento,
                "duracion": duracion_curso,
                "nombre_instructor": nombre_instructor,
                "ficha_instructor": ficha_instructor,
                "fecha_inicio": f_inicio,
                "fecha_termino": f_termino,
                "dia_inicio": dia_i,
                "mes_inicio": mes_i,
                "anio_inicio": anio_i,
                "dia_termino": dia_t,
                "mes_termino": mes_t,
                "anio_termino": anio_t,
                "categoria": cat,
                "categoria_trabajador": cat,
                "nivel": niv,
                "departamento": depto,
                "tipo_curso": tipo_curso,
                "nombre_supervisor": nombre_supervisor,
                "ficha_supervisor": ficha_supervisor
            }

            for p in plantillas_ind:
                t_path = os.path.join('templates', p)
                if os.path.exists(t_path):
                    if p.endswith('.docx'):
                        try:
                            doc = DocxTemplate(t_path)
                            doc.render(ctx_ind)
                            ficha_segura = re.sub(r'[^a-zA-Z0-9]', '', ficha)
                            out_name = f"{id_evento}_{ficha_segura}_{p.replace('.docx', '')}.docx"
                            doc.save(os.path.join(evento_dir, out_name))
                            docs_gen.append(out_name)
                        except Exception as ex:
                            print(f"Error en individual {p}: {ex}")
                    elif p.endswith('.xlsx'):
                        if p == 'SCPM-07.xlsx':
                            continue
                        try:
                            wb = openpyxl.load_workbook(t_path)
                            for ws in wb.worksheets:
                                for row_cells in ws.iter_rows():
                                    for cell in row_cells:
                                        if cell.value and isinstance(cell.value, str) and '{{' in cell.value:
                                            try:
                                                cell.value = render_template_string(cell.value, ctx_ind)
                                                ws.row_dimensions[cell.row].height = None
                                                if cell.alignment:
                                                    new_align = copy.copy(cell.alignment)
                                                    new_align.wrap_text = True
                                                    new_align.shrink_to_fit = False
                                                    cell.alignment = new_align
                                                else:
                                                    cell.alignment = openpyxl.styles.Alignment(wrap_text=True)
                                            except Exception:
                                                pass
                            ficha_segura = re.sub(r'[^a-zA-Z0-9]', '', ficha)
                            out_name = f"{id_evento}_{ficha_segura}_{p.replace('.xlsx', '')}.xlsx"
                            wb.save(os.path.join(evento_dir, out_name))
                            docs_gen.append(out_name)
                        except Exception as ex:
                            print(f"Error en individual xlsx {p}: {ex}")
                    elif p.endswith('.xls'):
                        try:
                            ficha_segura = re.sub(r'[^a-zA-Z0-9]', '', ficha)
                            out_name = f"{id_evento}_{ficha_segura}_{p}"
                            shutil.copy(t_path, os.path.join(evento_dir, out_name))
                            docs_gen.append(out_name)
                        except Exception as ex:
                            print(f"Error en individual xls {p}: {ex}")

        # 9. Generacion de formatos grupales del evento (1 documento por plantilla grupal)
        ctx_grp = {
            "clave_evento": id_evento,
            "nombre_evento": nombre_curso,
            "nombre_curso": nombre_curso,
            "duracion": duracion_curso,
            "nombre_instructor": nombre_instructor,
            "ficha_instructor": ficha_instructor,
            "nombre_supervisor": nombre_supervisor,
            "ficha_supervisor": ficha_supervisor,
            "fecha_inicio": f_inicio,
            "fecha_termino": f_termino,
            "dia_inicio": dia_i,
            "mes_inicio": mes_i,
            "anio_inicio": anio_i,
            "dia_termino": dia_t,
            "mes_termino": mes_t,
            "anio_termino": anio_t,
            "lista_participantes": lista_participantes,
            "tipo_curso": tipo_curso
        }

        # Filtra participantes activos (excluyendo bajas/no aptos previos) para el formato SCPM-07
        lista_activos = []
        for p_data in lista_participantes:
            existe = HistorialCapacitacion.query.filter_by(
                id_evento=id_evento,
                ficha_trabajador=p_data['ficha']
            ).first()
            if not existe or (existe.estado != 'BAJA' and existe.estado != 'NO APTO'):
                p_copy = p_data.copy()
                p_copy['calificacion'] = existe.calificacion if existe and existe.calificacion is not None else ""
                p_copy['nombre_trabajador'] = p_data['nombre_completo']
                lista_activos.append(p_copy)

        for p in plantillas_grp:
            if p == 'SCPM-07.xlsx':
                ctx_grp["lista_participantes"] = lista_activos
            else:
                ctx_grp["lista_participantes"] = lista_participantes

            t_path = os.path.join('templates', p)
            if os.path.exists(t_path):
                if p.endswith('.docx'):
                    try:
                        doc = DocxTemplate(t_path)
                        doc.render(ctx_grp)
                        out_name = f"{id_evento}_EVENTO_{p.replace('.docx', '')}.docx"
                        doc.save(os.path.join(evento_dir, out_name))
                        docs_gen.append(out_name)
                    except Exception as ex:
                        print(f"Error en grupal {p}: {ex}")
                elif p.endswith('.xlsx'):
                    try:
                        wb = openpyxl.load_workbook(t_path)
                        if p == 'SCPM-07.xlsx':
                            # Paginacion especial de 5 participantes por hoja para SCPM-07
                            original_ws = wb.worksheets[0]
                            original_title = original_ws.title
                            chunk_size = 5
                            parts = ctx_grp.get("lista_participantes", [])
                            chunks = [parts[i:i + chunk_size] for i in range(0, len(parts), chunk_size)]
                            if not chunks:
                                chunks = [[]]

                            worksheets = []
                            for chunk_idx in range(len(chunks)):
                                ws = wb.copy_worksheet(original_ws)
                                if len(chunks) > 1:
                                    ws.title = f"SCPM-07 ({chunk_idx + 1})"
                                else:
                                    ws.title = original_title + " (new)"
                                worksheets.append(ws)

                            wb.remove(original_ws)
                            if len(chunks) == 1:
                                worksheets[0].title = original_title

                            for chunk_idx, chunk in enumerate(chunks):
                                ws = worksheets[chunk_idx]

                                ficha_row, nombre_row = None, None
                                for row_cells in ws.iter_rows():
                                    for cell in row_cells:
                                        if cell.value and isinstance(cell.value, str):
                                            if '{% for participante in lista_participantes %}' in cell.value or 'participante.ficha' in cell.value:
                                                ficha_row = cell.row
                                            if 'paticipante.nombe_completo' in cell.value or 'participante.nombre_completo' in cell.value:
                                                nombre_row = cell.row

                                if ficha_row and nombre_row:
                                    for col in range(9, 14):
                                        ws.cell(row=ficha_row, column=col).value = ""
                                        ws.cell(row=nombre_row, column=col).value = ""
                                    for i, part in enumerate(chunk):
                                        col = 9 + i
                                        ws.cell(row=ficha_row, column=col).value = part.get('ficha', '')
                                        ws.cell(row=nombre_row, column=col).value = str(part.get('nombre_completo', '')).upper()

                                for row_cells in ws.iter_rows():
                                    for cell in row_cells:
                                        if cell.value and isinstance(cell.value, str) and '{' in cell.value:
                                            val = cell.value.replace('{% tr %}', '')
                                            val = val.replace('paticipante.nombe_completo', 'participante.nombre_completo')
                                            val = val.replace('{% for participante in lista_participantes %}', '')
                                            val = val.replace('{% endfor %}', '')
                                            try:
                                                cell.value = render_template_string(val, ctx_grp)
                                                ws.row_dimensions[cell.row].height = None
                                                if cell.alignment:
                                                    new_align = copy.copy(cell.alignment)
                                                    new_align.wrap_text = True
                                                    new_align.shrink_to_fit = False
                                                    cell.alignment = new_align
                                                else:
                                                    cell.alignment = openpyxl.styles.Alignment(wrap_text=True)
                                            except Exception:
                                                pass
                        else:
                            for ws in wb.worksheets:
                                for row_cells in ws.iter_rows():
                                    for cell in row_cells:
                                        if cell.value and isinstance(cell.value, str) and '{' in cell.value:
                                            try:
                                                val = cell.value.replace('{% tr %}', '')
                                                val = val.replace('paticipante.nombe_completo', 'participante.nombre_completo')
                                                cell.value = render_template_string(val, ctx_grp)
                                                ws.row_dimensions[cell.row].height = None
                                                if cell.alignment:
                                                    new_align = copy.copy(cell.alignment)
                                                    new_align.wrap_text = True
                                                    new_align.shrink_to_fit = False
                                                    cell.alignment = new_align
                                                else:
                                                    cell.alignment = openpyxl.styles.Alignment(wrap_text=True)
                                            except Exception:
                                                pass
                        out_name = f"{id_evento}_EVENTO_{p.replace('.xlsx', '')}.xlsx"
                        wb.save(os.path.join(evento_dir, out_name))
                        docs_gen.append(out_name)
                    except Exception as ex:
                        print(f"Error en grupal xlsx {p}: {ex}")
                elif p.endswith('.xlsm'):
                    # Poblamiento especial de tablas internas en SIRCE_Automatizado.xlsm (manteniendo macros VBA)
                    try:
                        wb = openpyxl.load_workbook(t_path, keep_vba=True)
                        modo_sirce = request.form.get('modo_sirce', 'con_datos')
                        if p == 'SIRCE_Automatizado.xlsm' and modo_sirce == 'con_datos':
                            if 'BD_TRABAJADORES' in wb.sheetnames:
                                ws_trab = wb['BD_TRABAJADORES']
                                header_row = 1
                                for r in range(1, 5):
                                    if any(c.value == 'CURP' for c in ws_trab[r]):
                                        header_row = r
                                        break
                                col_map_trab = {str(c.value).strip(): c.col_idx for c in ws_trab[header_row] if c.value}

                                # Se limpian las celdas sin usar delete_rows para no corromper las tablas ListObject de Excel
                                for r in range(header_row + 1, ws_trab.max_row + 1):
                                    for c in range(1, ws_trab.max_column + 1):
                                        ws_trab.cell(row=r, column=c).value = None

                                row_idx = header_row + 1
                                for part_data in trabajadores_curso.to_dict('records'):
                                    ficha = safe_str(part_data.get('FICHA', ''))
                                    curp = diccionario_curps.get(ficha, "")
                                    pat = str(part_data.get('PRIMER APELLIDO', '')).strip().title()
                                    mat = str(part_data.get('SEGUNDO APELLIDO', '')).strip().title()
                                    nom = str(part_data.get('NOMBRE', part_data.get('NOMBRE(S)', ''))).strip().title()

                                    if 'CURP' in col_map_trab:
                                        ws_trab.cell(row=row_idx, column=col_map_trab['CURP']).value = curp
                                    if 'NOMBRE' in col_map_trab:
                                        ws_trab.cell(row=row_idx, column=col_map_trab['NOMBRE']).value = nom
                                    if 'PRIMER APELLIDO' in col_map_trab:
                                        ws_trab.cell(row=row_idx, column=col_map_trab['PRIMER APELLIDO']).value = pat
                                    if 'SEGUNDO APELLIDO' in col_map_trab:
                                        ws_trab.cell(row=row_idx, column=col_map_trab['SEGUNDO APELLIDO']).value = mat
                                    elif 'SEGUNDI APELLICO' in col_map_trab:
                                        ws_trab.cell(row=row_idx, column=col_map_trab['SEGUNDI APELLICO']).value = mat
                                    if 'FICHA' in col_map_trab:
                                        ws_trab.cell(row=row_idx, column=col_map_trab['FICHA']).value = ficha
                                    row_idx += 1

                                from openpyxl.utils import get_column_letter
                                if ws_trab.tables:
                                    max_col_let = get_column_letter(ws_trab.max_column)
                                    end_r = row_idx - 1 if row_idx > header_row + 1 else header_row + 1
                                    for t in ws_trab.tables.values():
                                        t.ref = f"A{header_row}:{max_col_let}{end_r}"

                            if 'BD_CURSOS' in wb.sheetnames:
                                ws_cur = wb['BD_CURSOS']
                                header_row_cur = 1
                                for r in range(1, 5):
                                    if any(c.value == 'ID CURSO' for c in ws_cur[r]):
                                        header_row_cur = r
                                        break
                                col_map_cur = {str(c.value).strip(): c.col_idx for c in ws_cur[header_row_cur] if c.value}

                                for r in range(header_row_cur + 1, ws_cur.max_row + 1):
                                    for c in range(1, ws_cur.max_column + 1):
                                        ws_cur.cell(row=r, column=c).value = None

                                c_row = header_row_cur + 1
                                if 'ID CURSO' in col_map_cur:
                                    ws_cur.cell(row=c_row, column=col_map_cur['ID CURSO']).value = id_evento
                                if 'NOMBRE CURSO' in col_map_cur:
                                    ws_cur.cell(row=c_row, column=col_map_cur['NOMBRE CURSO']).value = nombre_curso
                                if 'DURACION' in col_map_cur:
                                    ws_cur.cell(row=c_row, column=col_map_cur['DURACION']).value = duracion_curso
                                if 'FEC INICIO' in col_map_cur:
                                    ws_cur.cell(row=c_row, column=col_map_cur['FEC INICIO']).value = f_inicio
                                if 'FEC TERMINO' in col_map_cur:
                                    ws_cur.cell(row=c_row, column=col_map_cur['FEC TERMINO']).value = f_termino

                        out_name = f"{id_evento}_EVENTO_{p.replace('.xlsm', '')}.xlsm"
                        wb.save(os.path.join(evento_dir, out_name))
                        docs_gen.append(out_name)
                    except Exception as ex:
                        print(f"Error en grupal xlsm {p}: {ex}")
                elif p.endswith('.xls'):
                    try:
                        out_name = f"{id_evento}_EVENTO_{p}"
                        shutil.copy(t_path, os.path.join(evento_dir, out_name))
                        docs_gen.append(out_name)
                    except Exception as ex:
                        print(f"Error en grupal xls {p}: {ex}")

        # 10. Registro o actualizacion del curso y sus participantes en la base de datos MySQL
        curso = Curso.query.filter_by(id_evento=id_evento).first()
        if not curso:
            curso = Curso(id_evento=id_evento, nombre_evento=nombre_curso, fase_actual=4, tipo_curso=tipo_curso)
            db.session.add(curso)
        else:
            curso.fase_actual = 4
            curso.nombre_evento = nombre_curso
            curso.tipo_curso = tipo_curso

        for part in historial:
            existe = HistorialCapacitacion.query.filter_by(
                id_evento=id_evento,
                ficha_trabajador=part['ficha_trabajador']
            ).first()
            if not existe:
                nuevo = HistorialCapacitacion(
                    id_evento=id_evento,
                    ficha_trabajador=part['ficha_trabajador'],
                    nombre_trabajador=part['nombre_trabajador']
                )
                db.session.add(nuevo)

        db.session.commit()

        return jsonify({
            "status": "success",
            "message": f"¡Éxito! Se generaron {len(docs_gen)} documentos.",
            "data": docs_gen
        }), 200

    except Exception as e:
        db.session.rollback()
        return error_response("Error del Motor: ", 500, e)


def download_docs(id_evento):
    """
    Endpoint `GET /api/download-docs/<id_evento>`.
    Convierte los documentos Word del evento a PDF mediante el servicio Gotenberg, genera un
    archivo unificado `<id_evento>_PDF_Maestro.pdf` y devuelve un paquete `.zip` clasificado
    en subcarpetas (`Cedulas/`, `Constancias/`, `Formatos SCPM/`, `SIRCE/`, `Instructor/`).
    """
    try:
        id_ev = str(id_evento).strip()
        if '..' in id_ev or '/' in id_ev or '\\' in id_ev:
            return jsonify({"status": "error", "message": "ID Inválido"}), 400

        evento_dir = event_dir(id_ev, create=True)
        if not os.path.exists(evento_dir):
            return jsonify({"status": "error", "message": "Sin documentos generados."}), 404

        all_files = os.listdir(evento_dir)
        if not all_files:
            return jsonify({"status": "error", "message": "No hay documentos."}), 404

        docx_files = [f for f in all_files if f.endswith('.docx')]
        xlsx_files = [f for f in all_files if f.endswith('.xlsx')]
        xlsm_files = [f for f in all_files if f.endswith('.xlsm')]
        xls_files = [f for f in all_files if f.endswith('.xls')]

        gotenberg_url = os.environ.get('GOTENBERG_URL', 'http://gotenberg:3000')

        # 1. Conversion de cada .docx a .pdf invocando la API REST de LibreOffice en Gotenberg
        for docx_name in docx_files:
            filepath = os.path.join(evento_dir, docx_name)
            with open(filepath, 'rb') as f:
                try:
                    res = requests.post(
                        f"{gotenberg_url}/forms/libreoffice/convert",
                        files={'files': (docx_name, f)},
                        timeout=300
                    )
                    if res.status_code == 200:
                        pdf_path = os.path.join(evento_dir, docx_name.replace('.docx', '.pdf'))
                        with open(pdf_path, 'wb') as pdf_file:
                            pdf_file.write(res.content)
                except Exception as e:
                    print(f"Error convirtiendo {docx_name} a PDF: {e}")

        # 2. Union de todos los PDFs individuales en un solo PDF Maestro para impresion rapida
        pdf_files = sorted([f for f in os.listdir(evento_dir) if f.endswith('.pdf')])
        merger = PdfWriter()
        for pdf in pdf_files:
            merger.append(os.path.join(evento_dir, pdf))

        m_path = os.path.join(evento_dir, f"{id_ev}_PDF_Maestro.pdf")
        if pdf_files:
            merger.write(m_path)
        merger.close()

        def get_folder(fname):
            """Determina la subcarpeta dentro del archivo ZIP segun el tipo de documento."""
            lower_name = fname.lower()
            if 'cédula' in lower_name or 'cedula' in lower_name:
                return 'Cedulas/'
            if 'dc-3' in lower_name or 'dc3' in lower_name or 'constancias' in lower_name:
                return 'Constancias/'
            if 'scpm' in lower_name:
                return 'Formatos SCPM/'
            if 'sirce' in lower_name:
                return 'SIRCE/'
            if 'compromiso' in lower_name or 'informe' in lower_name or 'instructor' in lower_name:
                return 'Instructor/'
            return 'Otros/'

        # 3. Construccion del archivo ZIP en memoria
        mem_file = io.BytesIO()
        with zipfile.ZipFile(mem_file, 'w', zipfile.ZIP_DEFLATED) as zf:
            if os.path.exists(m_path):
                zf.write(m_path, f"{id_ev}_PDF_Maestro.pdf")
            for f_list in [docx_files, xlsx_files, xlsm_files, xls_files]:
                for fname in f_list:
                    zf.write(os.path.join(evento_dir, fname), f"{get_folder(fname)}{fname}")
        mem_file.seek(0)

        return send_file(
            mem_file,
            mimetype='application/zip',
            as_attachment=True,
            download_name=f"Expediente_{id_ev}.zip"
        )
    except Exception as e:
        return error_response("Ocurrió un error interno", 500, e)


def extract_pdf():
    """
    Endpoint `POST /api/extract-pdf`.
    Analiza un archivo PDF de lista de asistencia (`SCPM-01`) y devuelve un archivo Excel
    estructurado (`Guia de Asistencia - <id_ev>.xlsx`) listo para revision y generacion documental.
    Incluye en los encabezados HTTP `X-Workers-Count` el total de participantes detectados.
    """
    if 'file' not in request.files:
        return jsonify({"status": "error", "message": "Sin archivo adjunto"}), 400
    file = request.files['file']
    if not file.filename.endswith('.pdf'):
        return jsonify({"status": "error", "message": "Solo .pdf"}), 400

    try:
        pdf_bytes = io.BytesIO(file.read())
        with pdfplumber.open(pdf_bytes) as pdf:
            texto_completo = ""
            for page in pdf.pages:
                texto_completo += page.extract_text(x_tolerance=2, y_tolerance=2) + "\n"

            txt_flat = re.sub(r'\s+', ' ', texto_completo.upper().replace('|', ' '))

            # 1. Extraccion del ID SIRHN del evento
            id_m = re.search(r"ID SIRHN GENERADO\s*(\d+)", txt_flat)
            id_ev = id_m.group(1) if id_m else "0"

            # 2. Busqueda del nombre del evento primero en celdas de tablas y luego por expresion regular
            nombre_ev = None
            for page in pdf.pages:
                for tabla in page.extract_tables():
                    for r_idx, fila in enumerate(tabla):
                        for c_idx, celda in enumerate(fila):
                            if celda and isinstance(celda, str):
                                c_up = celda.upper().strip()
                                c_up_flat = c_up.replace('\n', ' ')

                                prefixes = [
                                    'NOMBRE DEL EVENTO', 'NOMBRE DEL CURSO',
                                    'NOMBRE DEL TALLER', 'NOMBRE DEL PROYECTO'
                                ]

                                matched_prefix = None
                                for p in prefixes:
                                    if c_up_flat.startswith(p):
                                        matched_prefix = p
                                        break

                                if matched_prefix:
                                    if ':' in c_up_flat:
                                        val = c_up_flat.split(':', 1)[1].strip()
                                        if val:
                                            nombre_ev = val
                                    else:
                                        val = c_up_flat[len(matched_prefix):].strip()
                                        if val and val not in ('EVENTO', 'CURSO', 'TALLER', 'PROYECTO'):
                                            nombre_ev = val

                                    if not nombre_ev and c_idx + 1 < len(fila) and fila[c_idx + 1]:
                                        v = str(fila[c_idx + 1]).replace('\n', ' ').strip()
                                        if v and not v.upper().startswith('FECHA') and v.upper() not in ('EVENTO', 'CURSO', 'TALLER', 'PROYECTO'):
                                            nombre_ev = v
                                    if not nombre_ev and r_idx + 1 < len(tabla) and tabla[r_idx + 1][c_idx]:
                                        v = str(tabla[r_idx + 1][c_idx]).replace('\n', ' ').strip()
                                        if v and v.upper() not in ('EVENTO', 'CURSO', 'TALLER', 'PROYECTO'):
                                            nombre_ev = v
                        if nombre_ev:
                            break
                    if nombre_ev:
                        break

            if not nombre_ev or nombre_ev in ('EVENTO', 'CURSO', 'TALLER', 'PROYECTO'):
                nombre_ev = None
                stop_words = r"\b(?:FECHA|DURACI[OÓ]N|OBJETIVO|PERIODO|MODALIDAD|HORARIO|SEDE|LUGAR|INSTRUCTOR|ID SIRHN|ALCANCE|PERFIL|DIRIGIDO|TIPO|PARTICIPANTES|CLAVE|NO\.|VIGENCIA)\b"
                nom_m = re.search(
                    r"\b(?:NOMBRE DEL (?:EVENTO|CURSO|TALLER|PROYECTO)|ACCI[OÓ]N DE CAPACITACI[OÓ]N|TEMA|CURSO)\b\s*[:\-]?\s*(.+?)(?=\s+" + stop_words + r"|$)",
                    txt_flat
                )

                if not nom_m:
                    nom_m = re.search(
                        r"\b(?:NOMBRE DEL (?:EVENTO|CURSO|TALLER|PROYECTO)|ACCI[OÓ]N DE CAPACITACI[OÓ]N|TEMA|CURSO)\b\s*[:\-]?\s*([^\n]{1,150})",
                        texto_completo.upper()
                    )

                if nom_m:
                    val = nom_m.group(1).strip()
                    if val.startswith("CURSO CC"):
                        val = val.replace("CURSO CC", "").strip()
                    if val:
                        nombre_ev = val

            if not nombre_ev:
                nombre_ev = "SIN DATO"
            else:
                if nombre_ev.startswith("CURSO CC"):
                    nombre_ev = nombre_ev.replace("CURSO CC", "").strip()

            # 3. Extraccion de fecha de inicio, fecha de termino y duracion en horas
            f_matches = re.findall(r"FECHA DE INICIO\s*(\d{2}/\d{2}/\d{4})\s*(\d{2}/\d{2}/\d{4})\s*(\d+)", txt_flat)
            if not f_matches:
                f_matches = re.findall(r"(\d{2}/\d{2}/\d{4})\s*(\d{2}/\d{2}/\d{4})\s*(\d+)", txt_flat)
            if f_matches:
                f_ini, f_term, dur = f_matches[-1]
            else:
                f_ini, f_term, dur = "SIN DATO", "SIN DATO", "0"

            # 4. Extraccion de nombre y ficha del instructor y del supervisor que valida
            inst_m = re.search(
                r"HORARIO.*?\d{2}:\d{2}[^\s]*\s+([A-ZÑÁÉÍÓÚÜ\s\.]+?)\s*(?:F-?)?\s*(?:(?:DES\.\s*)?EJECUTIVO)?\s*(?:F-?)?\s*(\d{5,6})",
                txt_flat
            )
            if inst_m:
                raw_inst = inst_m.group(1).strip()
                n_inst = re.sub(r'\b(?:DES\.\s*)?EJECUTIVO\b', '', raw_inst).strip()
                n_inst = re.sub(r'\s+', ' ', n_inst)
                f_inst = inst_m.group(2).strip()
            else:
                n_inst = "SIN DATO"
                f_inst = "SIN DATO"

            val_m = re.search(
                r"ID SIRHN GENERADO\s*\d+\s+[A-ZÑÁÉÍÓÚÜ\.\s]+?\s*(?:F-?)?\s*\d{5,6}\s+([A-ZÑÁÉÍÓÚÜ\.\s]+?)\s*(?:F-?)?\s*(\d{5,6})",
                txt_flat
            )
            n_sup = val_m.group(1).strip() if val_m else "SIN DATO"
            f_sup = val_m.group(2).strip() if val_m else "SIN DATO"

            try:
                d_i, m_i, a_i = f_ini.split('/')
            except Exception:
                d_i, m_i, a_i = "", "", ""
            try:
                d_t, m_t, a_t = f_term.split('/')
            except Exception:
                d_t, m_t, a_t = "", "", ""

            def dividir_nombre(full_name):
                """
                Separa un nombre completo en formato 'APELLIDO_PATERNO APELLIDO_MATERNO NOMBRE(S)',
                respetando preposiciones compuestas ('DE', 'LA', 'DEL', 'LOS', etc.).
                """
                parts = str(full_name).strip().split()
                if not parts:
                    return "", "", ""
                prefixes = {"DE", "LA", "LAS", "DEL", "LOS", "MAC", "SAN", "SANTA", "Y"}
                g, c = [], ""
                for p in parts:
                    if p.upper() in prefixes:
                        c += p + " "
                    else:
                        c += p
                        g.append(c)
                        c = ""
                if len(g) == 1:
                    return g[0], "", ""
                elif len(g) == 2:
                    return g[0], "", g[1]
                else:
                    return g[0], g[1], " ".join(g[2:])

            # 5. Recorrido de las tablas de participantes en cada pagina del PDF
            participantes = []
            for page in pdf.pages:
                for tabla in page.extract_tables():
                    if tabla and len(tabla) > 0 and "Ficha" in str(tabla[0]):
                        for fila in tabla[1:]:
                            if len(fila) >= 5 and fila[1]:
                                ficha = str(fila[1]).replace('\n', '').strip()
                                if ficha.isdigit():
                                    pat, mat, nom = dividir_nombre(str(fila[2]).replace('\n', ' '))
                                    participantes.append({
                                        "ID": id_ev,
                                        "NOMBRE DEL EVENTO": nombre_ev,
                                        "FICHA": ficha,
                                        "PRIMER APELLIDO": pat,
                                        "SEGUNDO APELLIDO": mat,
                                        "NOMBRE": nom,
                                        "NIVEL": str(fila[3]).replace('\n', '').strip(),
                                        "CATEGORIA": str(fila[4]).replace('\n', ' ').strip(),
                                        "DIA I": d_i, "MES I": m_i, "AÑO I": a_i,
                                        "DIA T": d_t, "MES T": m_t, "AÑO T": a_t,
                                        "DURACION HORAS": dur,
                                        "NOMBRE INSTRUCTOR": n_inst.title(),
                                        "FICHA INSTRUCTOR": f_inst,
                                        "NOMBRE VALIDA (AREA USUARIA)": n_sup.title(),
                                        "FICHA VALIDA (AREA USUARIA)": f_sup
                                    })

        mem_file = io.BytesIO()
        with pd.ExcelWriter(mem_file, engine='openpyxl') as writer:
            if participantes:
                pd.DataFrame(participantes).to_excel(writer, sheet_name='Base_Datos', index=False)
            else:
                pd.DataFrame([{"Error": "No se encontraron trabajadores"}]).to_excel(writer, sheet_name='Errores')

        mem_file.seek(0)

        response = send_file(
            mem_file,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f"Guia de Asistencia - {id_ev}.xlsx"
        )

        # Cabecera personalizada leida por el frontend para mostrar cuantos trabajadores se extrajeron
        response.headers['X-Workers-Count'] = str(len(participantes)) if participantes else "0"
        response.headers['Access-Control-Expose-Headers'] = 'X-Workers-Count'

        return response
    except Exception as e:
        return error_response("Ocurrió un error interno", 500, e)


def generate_single_worker_docs(id_evento, worker_data):
    """
    Genera los formatos individuales (.docx / .xlsx / .xls) para un unico trabajador agregado
    manualmente a un evento que ya cuenta con `event_context.json`.
    """
    evento_dir = event_dir(str(id_evento).strip(), create=True)
    ctx_path = os.path.join(evento_dir, 'event_context.json')
    if not os.path.exists(ctx_path):
        return False

    with open(ctx_path, 'r', encoding='utf-8') as f:
        event_ctx = json.load(f)

    ficha = worker_data.get('ficha', '')
    nombre_completo = worker_data.get('nombre', '').strip().title()

    parts = str(nombre_completo).strip().split()
    if not parts:
        return False
    prefixes = {"DE", "LA", "LAS", "DEL", "LOS", "MAC", "SAN", "SANTA", "Y"}
    g, c = [], ""
    for p in parts:
        if p.upper() in prefixes:
            c += p + " "
        else:
            c += p
            g.append(c)
            c = ""
    if len(g) == 1:
        pat, mat, nom = g[0], "", ""
    elif len(g) == 2:
        pat, mat, nom = g[0], "", g[1]
    else:
        pat, mat, nom = g[0], g[1], " ".join(g[2:])

    curp = worker_data.get('curp', '')
    cat = worker_data.get('categoria', '')
    niv = worker_data.get('nivel', '')
    depto = worker_data.get('departamento', '')

    if not curp:
        ruta_curp = os.path.join('catalogos', 'CURP.xlsx')
        if os.path.exists(ruta_curp):
            try:
                df_curp = pd.read_excel(ruta_curp)
                df_curp.columns = df_curp.columns.str.upper().str.strip()
                if 'FICHA' in df_curp.columns and 'CURP' in df_curp.columns:
                    match = df_curp[df_curp['FICHA'].astype(str) == str(ficha)]
                    if not match.empty:
                        curp = str(match.iloc[0]['CURP']).upper()
            except Exception:
                pass

    ctx_ind = {
        "ficha": ficha,
        "ficha_trabajador": ficha,
        "apellido_paterno": pat,
        "apellido_materno": mat,
        "nombre": nom,
        "nombres": nom,
        "nombre_completo": nombre_completo,
        "nombre_trabajador": nombre_completo,
        "curp": curp,
        "CURP": curp,
        "nombre_evento": event_ctx.get('nombre_curso', ''),
        "nombre_curso": event_ctx.get('nombre_curso', ''),
        "clave_evento": id_evento,
        "duracion": event_ctx.get('duracion', ''),
        "nombre_instructor": event_ctx.get('nombre_instructor', ''),
        "ficha_instructor": event_ctx.get('ficha_instructor', ''),
        "fecha_inicio": event_ctx.get('fecha_inicio', ''),
        "fecha_termino": event_ctx.get('fecha_termino', ''),
        "dia_inicio": event_ctx.get('dia_inicio', ''),
        "mes_inicio": event_ctx.get('mes_inicio', ''),
        "anio_inicio": event_ctx.get('anio_inicio', ''),
        "dia_termino": event_ctx.get('dia_termino', ''),
        "mes_termino": event_ctx.get('mes_termino', ''),
        "anio_termino": event_ctx.get('anio_termino', ''),
        "categoria": cat,
        "categoria_trabajador": cat,
        "nivel": niv,
        "departamento": depto,
        "tipo_curso": event_ctx.get('tipo_curso', 'Actualización'),
        "nombre_supervisor": event_ctx.get('nombre_supervisor', ''),
        "ficha_supervisor": event_ctx.get('ficha_supervisor', '')
    }

    plantillas_ind = event_ctx.get('plantillas_ind', [])
    for p in plantillas_ind:
        t_path = os.path.join('templates', p)
        if os.path.exists(t_path):
            if p.endswith('.docx'):
                try:
                    doc = DocxTemplate(t_path)
                    doc.render(ctx_ind)
                    ficha_segura = re.sub(r'[^a-zA-Z0-9]', '', str(ficha))
                    out_name = f"{id_evento}_{ficha_segura}_{p.replace('.docx', '')}.docx"
                    doc.save(os.path.join(evento_dir, out_name))
                except Exception as ex:
                    print(f"Error en individual docx {p}: {ex}")
            elif p.endswith('.xlsx'):
                if p == 'SCPM-07.xlsx':
                    continue
                try:
                    wb = openpyxl.load_workbook(t_path)
                    for ws in wb.worksheets:
                        for row_cells in ws.iter_rows():
                            for cell in row_cells:
                                if cell.value and isinstance(cell.value, str) and '{{' in cell.value:
                                    try:
                                        cell.value = render_template_string(cell.value, ctx_ind)
                                        ws.row_dimensions[cell.row].height = None
                                        if cell.alignment:
                                            new_align = copy.copy(cell.alignment)
                                            new_align.wrap_text = True
                                            new_align.shrink_to_fit = False
                                            cell.alignment = new_align
                                        else:
                                            cell.alignment = openpyxl.styles.Alignment(wrap_text=True)
                                    except Exception:
                                        pass
                    ficha_segura = re.sub(r'[^a-zA-Z0-9]', '', str(ficha))
                    out_name = f"{id_evento}_{ficha_segura}_{p.replace('.xlsx', '')}.xlsx"
                    wb.save(os.path.join(evento_dir, out_name))
                except Exception as ex:
                    print(f"Error en individual xlsx {p}: {ex}")
            elif p.endswith('.xls'):
                try:
                    ficha_segura = re.sub(r'[^a-zA-Z0-9]', '', str(ficha))
                    out_name = f"{id_evento}_{ficha_segura}_{p}"
                    shutil.copy(t_path, os.path.join(evento_dir, out_name))
                except Exception:
                    pass
    return True
