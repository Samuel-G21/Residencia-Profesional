import docx

doc = docx.Document('backend/templates/SCPM-05 2025.docx')

for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            if 'Raymundo Acosta Chavarria' in cell.text:
                for p in cell.paragraphs:
                    if 'Raymundo Acosta Chavarria' in p.text:
                        p.text = p.text.replace('Ing. Raymundo Acosta Chavarria', '{{ nombre_supervisor }}')
                        p.text = p.text.replace('F- 282496', 'F- {{ ficha_supervisor }}')
                        p.text = p.text.replace('F-282496', 'F- {{ ficha_supervisor }}')

doc.save('backend/templates/SCPM-05 2025.docx')
