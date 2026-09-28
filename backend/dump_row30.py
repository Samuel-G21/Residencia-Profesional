import openpyxl

filepath = 'templates/SCPM-07.xlsx'
wb = openpyxl.load_workbook(filepath, keep_vba=False)
ws = wb.worksheets[0]

for col in range(1, ws.max_column + 1):
    cell = ws.cell(row=30, column=col)
    if cell.value:
        print(f"{cell.coordinate}: {cell.value!r}")
    cell2 = ws.cell(row=31, column=col)
    if cell2.value:
        print(f"{cell2.coordinate}: {cell2.value!r}")
