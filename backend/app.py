import os
import subprocess
import zipfile
import io
import shutil
import pandas as pd
import pymysql
import re
import pdfplumber
from contextlib import closing

from flask import send_file, Flask, request, jsonify
from flask_cors import CORS
from docxtpl import DocxTemplate
from pypdf import PdfWriter

app = Flask(__name__)
CORS(app)


# ==========================================
# 0. CATÁLOGOS STPS
# ==========================================
def cargar_catalogos():
    """Lee todos los Excel de STPS y los fusiona en RAM"""
    ruta_info = os.path.join('catalogos')
    os.makedirs(ruta_info, exist_ok=True)
    archivos = [f for f in os.listdir(ruta_info) if f.endswith('.xlsx') and not f.startswith('~$')]

    if not archivos:
        return pd.DataFrame()
    dfs = []
    for archivo in archivos:
        try:
            df_temp = pd.read_excel(os.path.join(ruta_info, archivo))
            df_temp.columns = df_temp.columns.str.upper().str.strip()
            dfs.append(df_temp)
        except Exception as e:
            print(f"Error leyendo catálogos {archivo}: {e}")

    if dfs:
        df_maestro = pd.concat(dfs, ignore_index=True)
        if 'FICHA' in df_maestro.columns:
            df_maestro['FICHA'] = df_maestro['FICHA'].astype(str).str.replace(r'\.0$', '', regex=True).str.strip()
        df_maestro = df_maestro.drop_duplicates(subset='FICHA', keep='last')
        return df_maestro
    return pd.DataFrame()


# ==========================================
# 1. CONEXIÓN A BASE DE DATOS
# ==========================================
def get_db_connection():
    try:
        conexion = pymysql.connect(
            host=os.getenv('DB_HOST', 'db'),
            user=os.getenv('DB_USER', 'pemex_user'),
            password=os.getenv('DB_PASSWORD', 'pemex_password'),
            database=os.getenv('DB_NAME', 'pemex_db'),
            cursorclass=pymysql.cursors.DictCursor
        )
        return conexion
    except Exception as e:
        print(f"Error conectando a BD: {e}")
        return None


@app.route('/', methods=['GET'])
def health_check():
    return jsonify({"status": "success", "message": "Sistema PEMEX Activo"}), 200


# ==========================================
# 2. CARGA Y LECTURA DE EXCEL
# ==========================================
@app.route('/api/upload-excel', methods=['POST'])
def upload_excel():
    if 'file' not in request.files:
        return jsonify({"status": "error", "message": "Sin archivo"}), 400

    file = request.files['file']
    if file.filename == '' or not file.filename.endswith('.xlsx'):
        return jsonify({"status": "error", "message": "Formato inválido (.xlsx)"}), 400

    try:
        df = pd.read_excel(file)
        df.columns = df.columns.str.upper().str.strip()
        col_id = 'ID'
        col_nom = 'NOMBRE DEL EVENTO'
        if col_id not in df.columns or col_nom not in df.columns:
            return jsonify({"status": "error", "message": "Falta ID o NOMBRE DEL EVENTO"}), 400
        cursos_df = df[[col_id, col_nom]].drop_duplicates()
        cursos_lista = cursos_df.rename(columns={col_id: 'id', col_nom: 'nombre'}).to_dict('records')
        return jsonify({"status": "success", "data": cursos_lista}), 200
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# 3. MOTOR DE INYECCIÓN DE DATOS (INDIVIDUALES Y GRUPALES)
# ==========================================
def safe_str(val):
    if pd.isna(val) or val == '': return ""
    try:
        f_val = float(val)
        if f_val.is_integer(): return str(int(f_val))
        return str(val).strip()
    except:
        return str(val).strip()


@app.route('/api/generate-docs', methods=['POST'])
def generate_docs():
    if 'file' not in request.files or 'id_evento' not in request.form:
        return jsonify({"status": "error", "message": "Faltan datos"}), 400

    file = request.files['file']
    id_evento = str(request.form['id_evento']).strip()

    try:
        df = pd.read_excel(file)
        df.columns = df.columns.str.upper().str.strip()
        df = df.fillna('')
        if 'ID' not in df.columns: return jsonify({"status": "error", "message": "Falta columna ID"}), 400
        df['ID'] = df['ID'].astype(str)
        trabajadores_curso = df[df['ID'] == id_evento]

        if trabajadores_curso.empty: return jsonify({"status": "error", "message": "Sin trabajadores"}), 404

        # Cargar Catálogo de CURPs desde la Base de Datos
        diccionario_curps = {}
        fichas_buscar = list(trabajadores_curso['FICHA'].dropna().unique())
        if fichas_buscar:
            conexion_curp = get_db_connection()
            if conexion_curp:
                with closing(conexion_curp):
                    with conexion_curp.cursor() as cursor:
                        format_strings = ','.join(['%s'] * len(fichas_buscar))
                        cursor.execute(f"SELECT ficha, curp FROM curps WHERE ficha IN ({format_strings})", tuple(fichas_buscar))
                        for row in cursor.fetchall():
                            diccionario_curps[row['ficha']] = row['curp'].upper()

        # Datos Generales del Evento
        row0 = trabajadores_curso.iloc[0]
        nombre_curso = str(row0.get('NOMBRE DEL EVENTO', '')).strip().upper()
        duracion_curso = str(row0.get('DURACION HORAS', row0.get('DURACION', ''))).strip()
        nombre_instructor = str(row0.get('INSTRUCTOR', row0.get('NOMBRE INSTRUCTOR', ''))).strip().title()
        ficha_instructor = str(row0.get('FICHA INSTRUCTOR', '')).strip()

        dia_i = safe_str(row0.get('DIA I', row0.get('dia i', ''))).zfill(2)
        mes_i = safe_str(row0.get('MES I', row0.get('mes i', ''))).zfill(2)
        anio_i = safe_str(row0.get('AÑO I', row0.get('año i', ''))).split('.')[0]
        dia_t = safe_str(row0.get('DIA T', row0.get('dia T', ''))).zfill(2)
        mes_t = safe_str(row0.get('MES T', row0.get('mes T', ''))).zfill(2)
        anio_t = safe_str(row0.get('AÑO T', row0.get('año T', ''))).split('.')[0]

        f_inicio = f"{dia_i}/{mes_i}/{anio_i}" if anio_i else ""
        f_termino = f"{dia_t}/{mes_t}/{anio_t}" if anio_t else ""

        # === CÓDIGO ORIGINAL (Comentado por si el cliente no aprueba la mejora) ===
        # plantillas_ind = [
        #     '1. Cédula registro actualizado 2025 COMBIANADA.docx',
        #     '2. Constancias de Habilidades DC-3 2026 COMBINADA.docx',
        #     'SCPM-04 COMBINADA.docx',
        #     'SCPM-04.docx',
        #     'SCPM-06 COMBINADA.docx'
        # ]
        # 
        # plantillas_grp = [
        #     '5. Carta Compromiso Instructor 2026 COMBINADA.docx',
        #     'FVC.docx',
        #     'Informe Técnico Instructor 2025.docx',
        #     'SCPM-05 2025.docx'
        # ]
        # ==========================================================================

        # === Selección de Documentos ===
        docs_seleccionados_str = request.form.get('docs_seleccionados')
        
        todas_ind = [
            '1. Cédula registro actualizado 2025 COMBIANADA.docx',
            '2. Constancias de Habilidades DC-3 2026 COMBINADA.docx',
            'SCPM-04 COMBINADA.docx',
            'SCPM-04.docx',
            'SCPM-06 COMBINADA.docx'
        ]
        
        todas_grp = [
            '5. Carta Compromiso Instructor 2026 COMBINADA.docx',
            'FVC.docx',
            'Informe Técnico Instructor 2025.docx',
            'SCPM-05 2025.docx',
            'SCPM-07.xlsx'
        ]

        if docs_seleccionados_str:
            lista_seleccionados = [d.strip() for d in docs_seleccionados_str.split(',')]
            plantillas_ind = [p for p in todas_ind if p in lista_seleccionados]
            plantillas_grp = [p for p in todas_grp if p in lista_seleccionados]
        else:
            plantillas_ind = todas_ind
            plantillas_grp = todas_grp
        # ==============================================

        evento_dir = os.path.join('outputs', id_evento)
        os.makedirs(evento_dir, exist_ok=True)
        for f in os.listdir(evento_dir):
            os.remove(os.path.join(evento_dir, f))

        docs_gen = []
        historial = []
        lista_participantes = []

        # 1. CICLO INDIVIDUAL (Por trabajador)
        for _, row in trabajadores_curso.iterrows():
            ficha = safe_str(row.get('FICHA', ''))
            nombres = str(row.get('NOMBRE', row.get('NOMBRE(S)', ''))).strip().title()
            pat = str(row.get('PRIMER APELLIDO', '')).strip().title()
            mat = str(row.get('SEGUNDO APELLIDO', '')).strip().title()
            nombre_completo = f"{pat} {mat} {nombres}".strip()
            cat = str(row.get('CATEGORIA', row.get('CATEGORÍA', ''))).strip()
            niv = str(row.get('NIVEL', '')).strip()
            depto = str(row.get('DEPARTAMENTO', '')).strip()
            curp = diccionario_curps.get(ficha, "SIN CURP EN CATALOGO")

            if ficha:
                lista_participantes.append({
                    "ficha": ficha,
                    "nombre_completo": nombre_completo,
                    "categoria": cat,
                    "nivel": niv,
                    "departamento": depto
                })
                historial.append((id_evento, ficha, nombre_completo))

            ctx_ind = {
                "ficha": ficha,
                "ficha_trabajador": ficha,
                "apellido_paterno": pat,
                "apellido_materno": mat,
                "nombre": nombres,
                "nombres": nombres,
                "nombre_completo": nombre_completo,
                "nombre_trabajador": nombre_completo,
                "curp": curp,
                "nombre_evento": nombre_curso,
                "nombre_curso": nombre_curso,
                "clave_evento": id_evento,
                "duracion": duracion_curso,
                "nombre_instructor": nombre_instructor,
                "ficha_instructor": ficha_instructor,
                "fecha_inicio": f_inicio,
                "fecha_termino": f_termino,
                "dia_inicio": dia_i,
                "mes_inicio": mes_i,
                "anio_inicio": anio_i,
                "dia_termino": dia_t,
                "mes_termino": mes_t,
                "anio_termino": anio_t,
                "categoria": cat,
                "categoria_trabajador": cat,
                "nivel": niv,
                "departamento": depto
            }

            for p in plantillas_ind:
                t_path = os.path.join('templates', p)
                if os.path.exists(t_path):
                    try:
                        doc = DocxTemplate(t_path)
                        doc.render(ctx_ind)
                        ficha_segura = re.sub(r'[^a-zA-Z0-9]', '', ficha)
                        out_name = f"{id_evento}_{ficha_segura}_{p.replace('.docx', '')}.docx"
                        doc.save(os.path.join(evento_dir, out_name))
                        docs_gen.append(out_name)
                    except Exception as ex:
                        print(f"Error en individual {p} (ficha {ficha}): {ex}")

        # 2. PROCESAMIENTO GRUPAL (Incluyendo SCPM-05 con lista_participantes)
        ctx_grp = {
            "clave_evento": id_evento,
            "nombre_evento": nombre_curso,
            "nombre_curso": nombre_curso,
            "duracion": duracion_curso,
            "nombre_instructor": nombre_instructor,
            "ficha_instructor": ficha_instructor,
            "fecha_inicio": f_inicio,
            "fecha_termino": f_termino,
            "dia_inicio": dia_i,
            "mes_inicio": mes_i,
            "anio_inicio": anio_i,
            "dia_termino": dia_t,
            "mes_termino": mes_t,
            "anio_termino": anio_t,
            "lista_participantes": lista_participantes
        }

        for p in plantillas_grp:
            t_path = os.path.join('templates', p)
            if os.path.exists(t_path):
                if p.endswith('.docx'):
                    try:
                        doc = DocxTemplate(t_path)
                        doc.render(ctx_grp)
                        out_name = f"{id_evento}_EVENTO_{p.replace('.docx', '')}.docx"
                        doc.save(os.path.join(evento_dir, out_name))
                        docs_gen.append(out_name)
                    except Exception as ex:
                        print(f"Error en grupal {p}: {ex}")
                elif p.endswith('.xlsx'):
                    try:
                        import openpyxl
                        wb = openpyxl.load_workbook(t_path)
                        
                        if 'Propuesta' in wb.sheetnames:
                            ws = wb['Propuesta']
                            
                            if ws['D9'].value and 'nombre_evento' in str(ws['D9'].value):
                                ws['D9'] = ctx_grp.get('nombre_evento', '')
                            
                            if ws['I54'].value and 'nombre_supervisor' in str(ws['I54'].value):
                                inst_name = ctx_grp.get('nombre_instructor', '')
                                inst_ficha = ctx_grp.get('ficha_instructor', '')
                                ws['I54'] = f"{inst_name} F-{inst_ficha}"

                            ws['I30'] = ''
                            ws['I31'] = ''
                            ws['M30'] = ''
                            
                            start_col = 9 
                            for idx, part in enumerate(ctx_grp.get('lista_participantes', [])):
                                ws.cell(row=30, column=start_col + idx).value = part.get('ficha', '')
                                ws.cell(row=31, column=start_col + idx).value = part.get('nombre_completo', '')
                        
                        out_name = f"{id_evento}_EVENTO_{p.replace('.xlsx', '')}.xlsx"
                        wb.save(os.path.join(evento_dir, out_name))
                        docs_gen.append(out_name)
                    except Exception as ex:
                        print(f"Error en grupal xlsx {p}: {ex}")

        # 3. GUARDADO EN BASE DE DATOS Y AVANCE DE FASE
        conexion = get_db_connection()
        if conexion:
            with closing(conexion):
                with conexion.cursor() as cursor:
                    cursor.execute("""
                        INSERT INTO cursos (id_evento, nombre_evento, fase_actual) 
                        VALUES (%s, %s, 4)
                        ON DUPLICATE KEY UPDATE fase_actual = 4, nombre_evento = %s
                    """, (id_evento, nombre_curso, nombre_curso))

                    if historial:
                        cursor.executemany("""
                            INSERT IGNORE INTO historial_capacitacion 
                            (id_evento, ficha_trabajador, nombre_trabajador) 
                            VALUES (%s, %s, %s)
                        """, historial)
                conexion.commit()

        return jsonify({
            "status": "success",
            "message": f"¡Éxito! Se generaron {len(docs_gen)} documentos (Individuales y Grupales).",
            "data": docs_gen
        }), 200

    except Exception as e:
        return jsonify({"status": "error", "message": f"Error del Motor: {str(e)}"}), 500


# ==========================================
# 4. COLA DE IMPRESIÓN (PDF & ZIP) CON GOTENBERG
# ==========================================
import requests

@app.route('/api/download-docs/<id_evento>', methods=['GET'])
def download_docs(id_evento):
    try:
        id_ev = str(id_evento).strip()
        # Protect against path traversal
        if '..' in id_ev or '/' in id_ev or '\\' in id_ev:
            return jsonify({"status": "error", "message": "ID Inválido"}), 400
            
        evento_dir = os.path.join('outputs', id_ev)
        if not os.path.exists(evento_dir):
            return jsonify({"status": "error", "message": "Sin documentos generados."}), 404

        docx_files = [f for f in os.listdir(evento_dir) if f.endswith('.docx')]
        if not docx_files:
            return jsonify({"status": "error", "message": "No hay documentos Word."}), 404

        gotenberg_url = os.getenv('GOTENBERG_URL', 'http://gotenberg:3000')

        # Convertir con Gotenberg uno por uno
        for docx in docx_files:
            filepath = os.path.join(evento_dir, docx)
            with open(filepath, 'rb') as f:
                try:
                    res = requests.post(f"{gotenberg_url}/forms/libreoffice/convert", files={'files': (docx, f)}, timeout=10)
                    if res.status_code == 200:
                        pdf_path = os.path.join(evento_dir, docx.replace('.docx', '.pdf'))
                        with open(pdf_path, 'wb') as pdf_file:
                            pdf_file.write(res.content)
                except Exception as e:
                    print(f"Error convirtiendo {docx} a PDF: {e}")

        # Unir PDFs
        pdf_files = sorted([f for f in os.listdir(evento_dir) if f.endswith('.pdf')])
        merger = PdfWriter()
        for pdf in pdf_files:
            merger.append(os.path.join(evento_dir, pdf))

        m_path = os.path.join(evento_dir, f"{id_ev}_PDF_Maestro.pdf")
        if pdf_files:
            merger.write(m_path)
        merger.close()

        # Comprimir
        mem_file = io.BytesIO()
        with zipfile.ZipFile(mem_file, 'w', zipfile.ZIP_DEFLATED) as zf:
            if os.path.exists(m_path):
                zf.write(m_path, f"{id_ev}_PDF_Maestro.pdf")
            for docx in docx_files:
                zf.write(os.path.join(evento_dir, docx), docx)
        mem_file.seek(0)

        # Se eliminó shutil.rmtree para permitir múltiples descargas
        return send_file(mem_file, mimetype='application/zip', as_attachment=True, download_name=f"Expediente_{id_ev}.zip")
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500



# ==========================================
# 5. MOTOR DE EXTRACCIÓN (SCPM-01)
# ==========================================
@app.route('/api/extract-pdf', methods=['POST'])
def extract_pdf():
    if 'file' not in request.files: return jsonify({"status": "error", "message": "Sin archivo adjunto"}), 400
    file = request.files['file']
    if not file.filename.endswith('.pdf'): return jsonify({"status": "error", "message": "Solo .pdf"}), 400

    try:
        pdf_bytes = io.BytesIO(file.read())
        with pdfplumber.open(pdf_bytes) as pdf:
            texto_completo = ""
            for page in pdf.pages: texto_completo += page.extract_text(x_tolerance=2, y_tolerance=2) + "\n"

            txt_flat = re.sub(r'\s+', ' ', texto_completo.upper().replace('|', ' '))

            id_m = re.search(r"ID SIRHN GENERADO\s*(\d+)", txt_flat)
            id_ev = id_m.group(1) if id_m else "0"

            nom_m = re.search(r"TIPO DE FORMACI[OÓ]N\s+(.*?)\s+(?:CURSO|EVENTO|TALLER)", txt_flat)
            nombre_ev = nom_m.group(1).strip() if nom_m else "SIN DATO"

            f_match = re.search(r"FECHA DE INICIO\s*(\d{2}/\d{2}/\d{4})\s*(\d{2}/\d{2}/\d{4})\s*(\d+)", txt_flat)
            if f_match:
                f_ini, f_term, dur = f_match.group(1), f_match.group(2), f_match.group(3)
            else:
                f_ini, f_term, dur = "SIN DATO", "SIN DATO", "0"

            inst_m = re.search(r"HORARIO.*?\d{2}:\d{2}.*?\s([A-ZÑ\s]+?)\s*F-?.*?(\d{5,6})", txt_flat)
            n_inst = inst_m.group(1).strip() if inst_m else "SIN DATO"
            f_inst = inst_m.group(2).strip() if inst_m else "SIN DATO"

            val_m = re.search(r"ID SIRHN GENERADO\s*\d+\s+[A-ZÑ\s]+?F-?\d{5,6}\s+([A-ZÑ\s]+?)\s*F-?(\d{5,6})", txt_flat)
            n_sup = val_m.group(1).strip() if val_m else "SIN DATO"
            f_sup = val_m.group(2).strip() if val_m else "SIN DATO"

            try:
                d_i, m_i, a_i = f_ini.split('/')
            except:
                d_i, m_i, a_i = "", "", ""
            try:
                d_t, m_t, a_t = f_term.split('/')
            except:
                d_t, m_t, a_t = "", "", ""

            def dividir_nombre(full_name):
                parts = str(full_name).strip().split()
                if not parts: return "", "", ""
                prefixes = {"DE", "LA", "LAS", "DEL", "LOS", "MAC", "SAN", "SANTA", "Y"}
                g, c = [], ""
                for p in parts:
                    if p.upper() in prefixes:
                        c += p + " "
                    else:
                        c += p;
                        g.append(c);
                        c = ""
                if len(g) == 1:
                    return g[0], "", ""
                elif len(g) == 2:
                    return g[0], "", g[1]
                else:
                    return g[0], g[1], " ".join(g[2:])

            participantes = []
            for page in pdf.pages:
                for tabla in page.extract_tables():
                    if tabla and len(tabla) > 0 and "Ficha" in str(tabla[0]):
                        for fila in tabla[1:]:
                            if len(fila) >= 5 and fila[1]:
                                ficha = str(fila[1]).replace('\n', '').strip()
                                if ficha.isdigit():
                                    pat, mat, nom = dividir_nombre(str(fila[2]).replace('\n', ' '))
                                    participantes.append({
                                        "ID": id_ev,
                                        "NOMBRE DEL EVENTO": nombre_ev,
                                        "FICHA": ficha,
                                        "PRIMER APELLIDO": pat, "SEGUNDO APELLIDO": mat, "NOMBRE": nom,
                                        "NIVEL": str(fila[3]).replace('\n', '').strip(),
                                        "CATEGORIA": str(fila[4]).replace('\n', ' ').strip(),
                                        "DIA I": d_i, "MES I": m_i, "AÑO I": a_i,
                                        "DIA T": d_t, "MES T": m_t, "AÑO T": a_t,
                                        "DURACION HORAS": dur,
                                        "NOMBRE INSTRUCTOR": n_inst.title(),
                                        "FICHA INSTRUCTOR": f_inst,
                                        "NOMBRE SUPERVISOR": n_sup.title(),
                                        "FICHA SUPERVISOR": f_sup
                                    })

        mem_file = io.BytesIO()
        with pd.ExcelWriter(mem_file, engine='openpyxl') as writer:
            if participantes:
                pd.DataFrame(participantes).to_excel(writer, sheet_name='Base_Datos', index=False)
            else:
                pd.DataFrame([{"Error": "No se encontraron trabajadores"}]).to_excel(writer, sheet_name='Errores')

        mem_file.seek(0)
        
        response = send_file(mem_file, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         as_attachment=True, download_name=f"Guia_{id_ev}.xlsx")
        
        # Expose custom header for frontend validation
        response.headers['X-Workers-Count'] = str(len(participantes)) if participantes else "0"
        response.headers['Access-Control-Expose-Headers'] = 'X-Workers-Count'
        
        return response
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# 6. DASHBOARD ESTADÍSTICO
# ==========================================
@app.route('/api/stats', methods=['GET'])
def get_stats():
    conexion = get_db_connection()
    if not conexion: return jsonify({"status": "error", "message": "BD desconectada"}), 500
    try:
        with closing(conexion):
            with conexion.cursor() as c:
                # 1. Total courses
                c.execute("SELECT COUNT(*) as total_cursos FROM cursos")
                tot_c = c.fetchone()['total_cursos']
                
                # 2. Total workers trained (all instances)
                c.execute("SELECT COUNT(*) as total_trabajadores FROM historial_capacitacion")
                tot_t = c.fetchone()['total_trabajadores']
                
                # 3. Unique workers trained
                c.execute("SELECT COUNT(DISTINCT ficha_trabajador) as unicos FROM historial_capacitacion")
                tot_unicos = c.fetchone()['unicos']
                
                # 4. Top 5 courses by participation
                c.execute(
                    "SELECT c.nombre_evento, COUNT(h.id) as total_capacitados FROM cursos c LEFT JOIN historial_capacitacion h ON c.id_evento = h.id_evento GROUP BY c.id_evento, c.nombre_evento ORDER BY total_capacitados DESC LIMIT 5")
                top = c.fetchall()
                
                # 5. Courses per phase
                c.execute(
                    "SELECT fase_actual, COUNT(*) as cantidad FROM cursos GROUP BY fase_actual ORDER BY fase_actual ASC"
                )
                fases_data = c.fetchall()
                
                # 6. Recent courses
                c.execute(
                    "SELECT nombre_evento, id_evento, fecha_registro, fase_actual FROM cursos ORDER BY fecha_registro DESC LIMIT 5"
                )
                recientes = c.fetchall()
                
        return jsonify({
            "status": "success", 
            "data": {
                "total_cursos": tot_c, 
                "total_trabajadores": tot_t, 
                "trabajadores_unicos": tot_unicos,
                "top_cursos": top,
                "cursos_por_fase": fases_data,
                "cursos_recientes": recientes
            }
        }), 200
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# 7. CATÁLOGOS OFICIALES
# ==========================================
@app.route('/api/update-catalog', methods=['POST'])
def update_catalog():
    if 'file' not in request.files or 'tipo' not in request.form:
        return jsonify({"status": "error", "message": "Falta archivo o tipo"}), 400
    file = request.files['file']
    tipo = str(request.form['tipo']).strip().lower()
    if not file.filename.endswith('.xlsx'): return jsonify({"status": "error", "message": "Solo Excel"}), 400
    nombres_cat = {
        "estados": "01 Cátalogo de Estados.xlsx", "municipios": "02 Cátalogo de Municipios.xlsx",
        "ocupaciones": "03 Cátalogo de Ocupaciones Pemex-STPS.xlsx",
        "escolaridad": "04 Cátalogo de Nivel de estudios (Escolaridad).xlsx",
        "probatorios": "05 Cátalogo de Documentos Probatorios.xlsx",
        "instituciones": "06 Cátalogo de Instituciones.xlsx",
        "areas_tematicas": "07 Cátalogo de Áreas Tematicas Pemex-STPS.xlsx",
        "agentes": "08 Cátalogo de Tipos de Agentes.xlsx",
        "modalidades": "09 Cátalogo de Modalidades de la Capacitacion.xlsx",
        "objetivos": "10 Cátalogo de Objetivos de la Capacitacion.xlsx", "curp": "CURP.xlsx"
    }
    if tipo not in nombres_cat: return jsonify({"status": "error", "message": "Desconocido"}), 400
    try:
        os.makedirs('catalogos', exist_ok=True)
        file.save(os.path.join('catalogos', nombres_cat[tipo]))
        return jsonify({"status": "success", "message": "Catálogo actualizado"}), 200
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# 7.5. HISTORIAL COMPLETO DE CURSOS
# ==========================================
@app.route('/api/cursos', methods=['GET'])
def get_all_cursos():
    conexion = get_db_connection()
    if not conexion: return jsonify({"status": "error", "message": "BD desconectada"}), 500
    try:
        with closing(conexion):
            with conexion.cursor() as c:
                c.execute("""
                    SELECT c.id_evento, c.nombre_evento, c.fase_actual, c.fecha_registro,
                           COUNT(h.id) as total_participantes
                    FROM cursos c
                    LEFT JOIN historial_capacitacion h ON c.id_evento = h.id_evento
                    GROUP BY c.id_evento, c.nombre_evento, c.fase_actual, c.fecha_registro
                    ORDER BY c.fecha_registro DESC
                """)
                cursos = c.fetchall()
        return jsonify({"status": "success", "data": cursos}), 200
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# ==========================================
# 8. CONSULTA DE FASE DEL EVENTO
# ==========================================
@app.route('/api/evento/<id_evento>', methods=['GET'])
def get_evento_fase(id_evento):
    conexion = get_db_connection()
    if not conexion:
        return jsonify({"status": "error", "message": "BD desconectada"}), 500

    try:
        with closing(conexion):
            with conexion.cursor() as cursor:
                cursor.execute("SELECT id_evento, nombre_evento, fase_actual FROM cursos WHERE id_evento = %s",
                               (id_evento,))
                evento = cursor.fetchone()

        if evento:
            return jsonify({"status": "success", "data": evento}), 200
        else:
            return jsonify({"status": "not_found", "message": "Evento nuevo", "fase_sugerida": 1}), 200

    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route('/api/evento/<id_evento>/fase', methods=['PUT'])
def update_evento_fase(id_evento):
    data = request.json
    nueva_fase = data.get('fase')
    if not nueva_fase:
        return jsonify({"status": "error", "message": "Falta la nueva fase"}), 400

    conexion = get_db_connection()
    if not conexion:
        return jsonify({"status": "error", "message": "BD desconectada"}), 500

    try:
        with closing(conexion):
            with conexion.cursor() as cursor:
                cursor.execute("UPDATE cursos SET fase_actual = %s WHERE id_evento = %s", (nueva_fase, id_evento))
                if cursor.rowcount == 0:
                    cursor.execute("INSERT INTO cursos (id_evento, nombre_evento, fase_actual) VALUES (%s, %s, %s)", (id_evento, 'Evento Manual', nueva_fase))
            conexion.commit()
        return jsonify({"status": "success", "message": f"Fase actualizada a {nueva_fase}"}), 200
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route('/api/evento/<id_evento>', methods=['DELETE'])
def delete_evento(id_evento):
    conexion = get_db_connection()
    if not conexion:
        return jsonify({"status": "error", "message": "BD desconectada"}), 500

    try:
        with closing(conexion):
            with conexion.cursor() as cursor:
                # La tabla historial_capacitacion tiene ON DELETE CASCADE según init.sql
                cursor.execute("DELETE FROM cursos WHERE id_evento = %s", (id_evento,))
                affected = cursor.rowcount
            conexion.commit()

        # También borrar archivos si existen
        id_ev = str(id_evento).strip()
        evento_dir = os.path.join('outputs', id_ev)
        if os.path.exists(evento_dir):
            import shutil
            shutil.rmtree(evento_dir, ignore_errors=True)

        if affected > 0:
            return jsonify({"status": "success", "message": "Evento eliminado"}), 200
        else:
            return jsonify({"status": "error", "message": "Evento no encontrado"}), 404
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)