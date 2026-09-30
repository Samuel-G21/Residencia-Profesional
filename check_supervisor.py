import os
import docx

templates = [
    'SCPM-05 2025.docx',
    'Informe Técnico Instructor 2025.docx',
    'SCPM-04 COMBINADA.docx',
    'SCPM-04.docx',
    'SCPM-06 COMBINADA.docx'
]

for t in templates:
    path = os.path.join('backend/templates', t)
    if not os.path.exists(path):
        print(f"Missing: {t}")
        continue
    
    doc = docx.Document(path)
    found = False
    for para in doc.paragraphs:
        if 'nombre_supervisor' in para.text:
            found = True
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if 'nombre_supervisor' in cell.text:
                    found = True
    print(f"{t}: {'YES' if found else 'NO'}")
