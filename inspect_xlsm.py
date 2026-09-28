import openpyxl

wb = openpyxl.load_workbook("D:\\BASE PLANTILLA SIRCE_Automatizado.xlsm", keep_vba=True)
print("Sheetnames:", wb.sheetnames)

if 'BD_TRABAJADORES' in wb.sheetnames:
    ws = wb['BD_TRABAJADORES']
    for r in range(1, 5):
        vals = [c.value for c in ws[r] if c.value]
        if vals:
            print(f"BD_TRABAJADORES Row {r}: {vals}")

if 'BD_CURSOS' in wb.sheetnames:
    ws = wb['BD_CURSOS']
    for r in range(1, 5):
        vals = [c.value for c in ws[r] if c.value]
        if vals:
            print(f"BD_CURSOS Row {r}: {vals}")
